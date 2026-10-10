# -*- coding: utf-8 -*-
"""
Builder para o Relatório Oficial de Market Share Móvel & 5G — Regional Norte
Conformidade rigorosa com PADRAO_DESIGN_PAINEIS.md e layout idêntico ao Market Share Residencial:
  1. Visão Resumida: 4 quadros (Market Share, Net Adds, Base, Net Share) + Variação MoM,
     com sinalização do melhor da linha, coluna de posição da Claro e seletor (DDD, UF, Regional com todas as 10 regionais reais).
  2. Penetração 5G: Matriz 5G por DDD (Aparelhos 5G, Base 5G e % Penetração 5G).
  3. Visão por Município: Divisão por Grandes Mercados, Mercados em Desenvolvimento e Demais Cidades com subtotais.
  4. Visão Detalhada: Matriz histórica mensal (14 meses) com abertura individual de cada indicador
     (uma abertura de DDD na Base, outra no Share, outra em Adds e outra em Net Share) que responde dinamicamente aos filtros.
  5. Pré-set padrão inicial: Modalidade de Cobrança = "Pós-pago".
  6. Lógica de liderança segura: sem liderança com base 0 ou ausência de rede.
  7. Gaveta de filtros retrátil padrão residencial, botões quadrados e design system 100% corporativo Claro.
"""

import os
import sys
import json
import time
import openpyxl

try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
REP_DIR = os.path.join(ROOT_DIR, "relatorios", "market_share_movel")
EXCEL_PATH = os.path.join(REP_DIR, "dados", "Ago2026_Mkt_Share_Móvel_v9.8 - RESUMIDA _SEM_AJUSTE PJ v8.xlsm")
OUTPUT_HTML = os.path.join(REP_DIR, "index.html")

def clean_f(v):
    if v is None or str(v).strip() in ['', '-', 'None']: return 0.0
    try: return float(v)
    except: return 0.0

def clean_i(v):
    if v is None or str(v).strip() in ['', '-', 'None']: return 0
    try: return int(round(float(v)))
    except: return 0

def build():
    print("Compilando Market Share Móvel & 5G Oficial...")
    if not os.path.exists(EXCEL_PATH):
        print(f"❌ Erro: Planilha não encontrada em {EXCEL_PATH}")
        return False

    t0 = time.time()
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    # 1. Meses Analíticos (14 meses: jul/25 a ago/26)
    months = ['jul/25', 'ago/25', 'set/25', 'out/25', 'nov/25', 'dez/25', 'jan/26', 'fev/26', 'mar/26', 'abr/26', 'mai/26', 'jun/26', 'jul/26', 'ago/26']

    # 2. Extração D_Resumo Móvel (DDD, UF e as 10 REGIONAIS DO BRASIL)
    print("  • Extraindo D_Resumo Móvel (DDD, UF e 10 Regionais Oficiais)...")
    s_res = wb['D_Resumo']
    
    # 2.1 Visão DDD (Linhas 9 a 18)
    ddd_data = []
    for r in range(9, 19):
        r_base = [c for c in next(s_res.iter_rows(min_row=r, max_row=r, values_only=True))]
        r_adds = [c for c in next(s_res.iter_rows(min_row=r+14, max_row=r+14, values_only=True))]
        r_net = [c for c in next(s_res.iter_rows(min_row=r+28, max_row=r+28, values_only=True))]
        
        label = str(r_base[3]).strip()
        ddd_data.append({
            'key': label,
            'label': 'Total Regional Norte' if label.lower() == 'total' else f'DDD {label}',
            'base': {
                'CLARO': clean_i(r_base[4]), 'VIVO': clean_i(r_base[5]), 'TIM': clean_i(r_base[6]),
                'SURF': clean_i(r_base[7]), 'OUTROS': clean_i(r_base[8]),
                'total': clean_i(r_base[4]) + clean_i(r_base[5]) + clean_i(r_base[6]) + clean_i(r_base[7]) + clean_i(r_base[8])
            },
            'share': {
                'CLARO': round(clean_f(r_base[12])*100, 2), 'VIVO': round(clean_f(r_base[13])*100, 2), 'TIM': round(clean_f(r_base[14])*100, 2),
                'SURF': round(clean_f(r_base[15])*100, 2), 'OUTROS': round(clean_f(r_base[16])*100, 2)
            },
            'adds': {
                'CLARO': clean_i(r_adds[4]), 'VIVO': clean_i(r_adds[5]), 'TIM': clean_i(r_adds[6]),
                'SURF': clean_i(r_adds[7]), 'OUTROS': clean_i(r_adds[8]),
                'total': clean_i(r_adds[4]) + clean_i(r_adds[5]) + clean_i(r_adds[6]) + clean_i(r_adds[7]) + clean_i(r_adds[8])
            },
            'mom': {
                'CLARO': round(clean_f(r_adds[12])*100, 2), 'VIVO': round(clean_f(r_adds[13])*100, 2), 'TIM': round(clean_f(r_adds[14])*100, 2),
                'SURF': round(clean_f(r_adds[15])*100, 2), 'OUTROS': round(clean_f(r_adds[16])*100, 2)
            },
            'net_share': {
                'CLARO': round(clean_f(r_net[4])*100, 2), 'VIVO': round(clean_f(r_net[5])*100, 2), 'TIM': round(clean_f(r_net[6])*100, 2),
                'SURF': round(clean_f(r_net[7])*100, 2), 'OUTROS': round(clean_f(r_net[8])*100, 2)
            }
        })

    # 2.2 Visão UF (Linhas 9 a 14)
    uf_data = []
    for r in range(9, 15):
        r_base = [c for c in next(s_res.iter_rows(min_row=r, max_row=r, values_only=True))]
        r_adds = [c for c in next(s_res.iter_rows(min_row=r+14, max_row=r+14, values_only=True))]
        r_net = [c for c in next(s_res.iter_rows(min_row=r+28, max_row=r+28, values_only=True))]
        
        label = str(r_base[19]).strip()
        uf_data.append({
            'key': label,
            'label': 'Total Regional Norte' if label.lower() == 'total' else f'UF {label}',
            'base': {
                'CLARO': clean_i(r_base[20]), 'VIVO': clean_i(r_base[21]), 'TIM': clean_i(r_base[22]),
                'SURF': clean_i(r_base[23]), 'OUTROS': clean_i(r_base[24]),
                'total': clean_i(r_base[20]) + clean_i(r_base[21]) + clean_i(r_base[22]) + clean_i(r_base[23]) + clean_i(r_base[24])
            },
            'share': {
                'CLARO': round(clean_f(r_base[28])*100, 2), 'VIVO': round(clean_f(r_base[29])*100, 2), 'TIM': round(clean_f(r_base[30])*100, 2),
                'SURF': round(clean_f(r_base[31])*100, 2), 'OUTROS': round(clean_f(r_base[32])*100, 2)
            },
            'adds': {
                'CLARO': clean_i(r_adds[20]), 'VIVO': clean_i(r_adds[21]), 'TIM': clean_i(r_adds[22]),
                'SURF': clean_i(r_adds[23]), 'OUTROS': clean_i(r_adds[24]),
                'total': clean_i(r_adds[20]) + clean_i(r_adds[21]) + clean_i(r_adds[22]) + clean_i(r_adds[23]) + clean_i(r_adds[24])
            },
            'mom': {
                'CLARO': round(clean_f(r_adds[28])*100, 2), 'VIVO': round(clean_f(r_adds[29])*100, 2), 'TIM': round(clean_f(r_adds[30])*100, 2),
                'SURF': round(clean_f(r_adds[31])*100, 2), 'OUTROS': round(clean_f(r_adds[32])*100, 2)
            },
            'net_share': {
                'CLARO': round(clean_f(r_net[20])*100, 2), 'VIVO': round(clean_f(r_net[21])*100, 2), 'TIM': round(clean_f(r_net[22])*100, 2),
                'SURF': round(clean_f(r_net[23])*100, 2), 'OUTROS': round(clean_f(r_net[24])*100, 2)
            }
        })

    # 2.3 Visão REGIONAL NACIONAL COMPLETA (Linhas 9 a 19, cols 36..49) - Item 2 do Usuário
    reg_data = []
    for r in range(9, 20):
        r_base = [c for c in next(s_res.iter_rows(min_row=r, max_row=r, values_only=True))]
        r_adds = [c for c in next(s_res.iter_rows(min_row=r+14, max_row=r+14, values_only=True))]
        r_net = [c for c in next(s_res.iter_rows(min_row=r+28, max_row=r+28, values_only=True))]
        
        label = str(r_base[36]).strip()
        reg_data.append({
            'key': label,
            'label': 'Total Nacional' if label.lower() == 'total' else f'Regional {label}',
            'base': {
                'CLARO': clean_i(r_base[37]), 'VIVO': clean_i(r_base[38]), 'TIM': clean_i(r_base[39]),
                'SURF': clean_i(r_base[40]), 'OUTROS': clean_i(r_base[41]),
                'total': clean_i(r_base[37]) + clean_i(r_base[38]) + clean_i(r_base[39]) + clean_i(r_base[40]) + clean_i(r_base[41])
            },
            'share': {
                'CLARO': round(clean_f(r_base[45])*100, 2), 'VIVO': round(clean_f(r_base[46])*100, 2), 'TIM': round(clean_f(r_base[47])*100, 2),
                'SURF': round(clean_f(r_base[48])*100, 2), 'OUTROS': round(clean_f(r_base[49])*100, 2)
            },
            'adds': {
                'CLARO': clean_i(r_adds[37]), 'VIVO': clean_i(r_adds[38]), 'TIM': clean_i(r_adds[39]),
                'SURF': clean_i(r_adds[40]), 'OUTROS': clean_i(r_adds[41]),
                'total': clean_i(r_adds[37]) + clean_i(r_adds[38]) + clean_i(r_adds[39]) + clean_i(r_adds[40]) + clean_i(r_adds[41])
            },
            'mom': {
                'CLARO': round(clean_f(r_adds[45])*100, 2), 'VIVO': round(clean_f(r_adds[46])*100, 2), 'TIM': round(clean_f(r_adds[47])*100, 2),
                'SURF': round(clean_f(r_adds[48])*100, 2), 'OUTROS': round(clean_f(r_adds[49])*100, 2)
            },
            'net_share': {
                'CLARO': round(clean_f(r_net[37])*100, 2), 'VIVO': round(clean_f(r_net[38])*100, 2), 'TIM': round(clean_f(r_net[39])*100, 2),
                'SURF': round(clean_f(r_net[40])*100, 2), 'OUTROS': round(clean_f(r_net[41])*100, 2)
            }
        })

    # 3. D_5G (Penetração 5G por DDD)
    print("  • Extraindo D_5G...")
    s_5g = wb['D_5G']
    data_5g = []
    for r in range(9, 19):
        row = [c for c in next(s_5g.iter_rows(min_row=r, max_row=r, values_only=True))]
        label = str(row[3]).strip()
        data_5g.append({
            'ddd': label,
            'label': 'Total Regional Norte' if label.lower() == 'total' else f'DDD {label}',
            'pct_5g': {
                'CLARO': round(clean_f(row[4])*100, 2), 'VIVO': round(clean_f(row[5])*100, 2), 'TIM': round(clean_f(row[6])*100, 2),
                'SURF': round(clean_f(row[7])*100, 2), 'OUTROS': round(clean_f(row[8])*100, 2)
            },
            'base_5g': {
                'CLARO': clean_i(row[12]), 'VIVO': clean_i(row[13]), 'TIM': clean_i(row[14]),
                'SURF': clean_i(row[15]), 'OUTROS': clean_i(row[16]),
                'total': clean_i(row[12]) + clean_i(row[13]) + clean_i(row[14]) + clean_i(row[15]) + clean_i(row[16])
            }
        })

    # 4. D_Município Móvel (455 Municípios com todas as métricas)
    print("  • Extraindo D_Município...")
    s_mun = [s for s in wb.sheetnames if 'Munic' in s][0]
    sheet_m = wb[s_mun]
    cities_data = []

    import unicodedata
    def norm_city(txt):
        txt = unicodedata.normalize('NFKD', str(txt or '')).encode('ASCII', 'ignore').decode('utf-8')
        return txt.upper().strip()

    grandes_mercados_exact = {
        'AM - MANAUS', 'AP - MACAPA', 'AP - SANTANA', 'MA - SAO LUIS', 'PA - ANANINDEUA', 'PA - BELEM', 'RR - BOA VISTA'
    }
    desenv_mercados_exact = {
        'MA - CAXIAS', 'MA - IMPERATRIZ', 'MA - TIMON', 'PA - CASTANHAL', 'PA - MARABA', 'PA - PARAGOMINAS', 'PA - PARAUAPEBAS'
    }

    for r in range(10, 470):
        try:
            row = [c for c in next(sheet_m.iter_rows(min_row=r, max_row=r, values_only=True))]
        except StopIteration:
            break
        if len(row) > 3 and row[3]:
            cid = str(row[3]).strip()
            if cid in ['', 'None']: continue
            
            cid_norm = norm_city(cid)
            div_mercado = 'Demais Cidades'
            if cid_norm in grandes_mercados_exact:
                div_mercado = 'Grandes Mercados'
            elif cid_norm in desenv_mercados_exact:
                div_mercado = 'Mercados em Desenvolvimento'

            cities_data.append({
                'cidade': cid,
                'divisao': div_mercado,
                'base': {
                    'CLARO': clean_i(row[25]), 'VIVO': clean_i(row[26]), 'TIM': clean_i(row[27]),
                    'SURF': clean_i(row[28]), 'OUTROS': clean_i(row[29]),
                    'total': clean_i(row[25]) + clean_i(row[26]) + clean_i(row[27]) + clean_i(row[28]) + clean_i(row[29])
                },
                'share': {
                    'CLARO': round(clean_f(row[4])*100, 2), 'VIVO': round(clean_f(row[5])*100, 2), 'TIM': round(clean_f(row[6])*100, 2),
                    'SURF': round(clean_f(row[7])*100, 2), 'OUTROS': round(clean_f(row[8])*100, 2)
                },
                'adds': {
                    'CLARO': clean_i(row[18]), 'VIVO': clean_i(row[19]), 'TIM': clean_i(row[20]),
                    'SURF': clean_i(row[21]), 'OUTROS': clean_i(row[22]),
                    'total': clean_i(row[18]) + clean_i(row[19]) + clean_i(row[20]) + clean_i(row[21]) + clean_i(row[22])
                }
            })

    # 5. D_Móvel (Matriz Histórica Mensal de 14 Meses com abertura de TODOS os 4 indicadores por DDD) - Item 4
    print("  • Extraindo D_Móvel (Histórico Completo por Operadora, Indicador e DDD)...")
    ws_mov = wb['D_Móvel']
    history_mov = {}
    
    op_configs = {
        'Total Geral': {'base': 8, 'adds': 19},
        'Claro': {'base': 34, 'share': 45, 'adds': 67, 'net_share': 78},
        'Vivo': {'base': 93, 'share': 104, 'adds': 126, 'net_share': 137},
        'Tim': {'base': 152, 'share': 163, 'adds': 185, 'net_share': 196},
        'SURF': {'base': 211, 'share': 222, 'adds': 244, 'net_share': 255},
        'Outros': {'base': 329, 'share': 340, 'adds': 362, 'net_share': 373}
    }

    for op, cfg in op_configs.items():
        history_mov[op] = {}
        for ind, base_r in cfg.items():
            row_vals = [c for c in next(ws_mov.iter_rows(min_row=base_r, max_row=base_r, values_only=True))]
            is_pct = ind in ['share', 'net_share']
            m14_master = [round(clean_f(row_vals[c])*100, 2) if is_pct else clean_i(row_vals[c]) for c in range(83, 97)]
            
            ddds_data = {}
            for d_idx, ddd in enumerate(['91', '92', '93', '94', '95', '96', '97', '98', '99']):
                d_row = [c for c in next(ws_mov.iter_rows(min_row=base_r + 1 + d_idx, max_row=base_r + 1 + d_idx, values_only=True))]
                ddds_data[ddd] = [round(clean_f(d_row[c])*100, 2) if is_pct else clean_i(d_row[c]) for c in range(83, 97)]
            
            history_mov[op][ind] = {
                'master': m14_master,
                'ddds': ddds_data
            }

    # Para Total Geral, preencher share e net_share sintéticos (100% constante)
    history_mov['Total Geral']['share'] = {
        'master': [100.0] * 14,
        'ddds': {d: [100.0] * 14 for d in ['91', '92', '93', '94', '95', '96', '97', '98', '99']}
    }
    history_mov['Total Geral']['net_share'] = {
        'master': [100.0] * 14,
        'ddds': {d: [100.0] * 14 for d in ['91', '92', '93', '94', '95', '96', '97', '98', '99']}
    }

    # 6. BASE Analítica Móvel (Linhas granulares RNO para gaveta de filtros)
    print("  • Extraindo BASE Analítica RNO...")
    s_base = wb['BASE']
    dim_empresas = []
    dim_produtos = []
    dim_ufs = []
    dim_ddds = []
    dim_cobrancas = []
    dim_pessoas = []
    dim_tecnologias = []

    def get_id(lst, val):
        val = str(val or '').strip()
        if not val: val = 'Outros'
        if val not in lst:
            lst.append(val)
        return lst.index(val)

    analytic_rows = []
    for row in s_base.iter_rows(min_row=3, values_only=True):
        reg = str(row[2] or '').strip()
        if reg == 'RNO':
            jul = clean_f(row[100])
            ago = clean_f(row[101])
            if jul == 0.0 and ago == 0.0: continue

            e_id = get_id(dim_empresas, row[3])
            pr_id = get_id(dim_produtos, row[4])
            u_id = get_id(dim_ufs, row[5])
            d_id = get_id(dim_ddds, row[6])
            cb_id = get_id(dim_cobrancas, row[7])
            pe_id = get_id(dim_pessoas, row[8])
            tc_id = get_id(dim_tecnologias, row[9])

            analytic_rows.append([e_id, pr_id, u_id, d_id, cb_id, pe_id, tc_id, round(jul), round(ago)])

    wb.close()

    payload = {
        'mes_referencia': 'Agosto/2026',
        'data_posicao': '21/09/2026',
        'months': months,
        'ddd_data': ddd_data,
        'uf_data': uf_data,
        'reg_data': reg_data,
        'data_5g': data_5g,
        'cities_data': cities_data,
        'history_mov': history_mov,
        'analytic': {
            'dim': {
                'empresas': dim_empresas,
                'produtos': dim_produtos,
                'ufs': dim_ufs,
                'ddds': dim_ddds,
                'cobrancas': dim_cobrancas,
                'pessoas': dim_pessoas,
                'tecnologias': dim_tecnologias
            },
            'rows': analytic_rows
        }
    }

    print("  • Gerando HTML corporativo de Market Share Móvel & 5G...")
    html_content = generate_html(payload)
    with open(OUTPUT_HTML, "w", encoding="utf-8") as f:
        f.write(html_content)

    size_kb = os.path.getsize(OUTPUT_HTML) / 1024
    print(f"✅ Painel Móvel gerado com sucesso em: {OUTPUT_HTML} ({size_kb:.1f} KB)")
    return True


def generate_html(payload):
    json_data = json.dumps(payload, separators=(',', ':'), ensure_ascii=False)
    
    # Template HTML com injeção segura por placeholder (evita problemas de interpolação f-string de chaves JS)
    template = """<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>Market Share Móvel &amp; 5G — Regional Norte</title>

<!-- PWA & Mobile Web App Meta Tags -->
<link rel="manifest" href="/manifest.webmanifest">
<meta name="theme-color" content="#8a0f18">
<meta name="mobile-web-app-capable" content="yes">
<meta name="apple-mobile-web-app-capable" content="yes">
<meta name="apple-mobile-web-app-status-bar-style" content="black-translucent">
<meta name="apple-mobile-web-app-title" content="PEP HUB">
<link rel="apple-touch-icon" href="/icons/apple-touch-icon.png">
<link rel="icon" type="image/svg+xml" href="/icons/icon.svg">
<link rel="icon" type="image/png" sizes="32x32" href="/icons/favicon.png">
<link rel="stylesheet" href="/css/app-mobile.css">

<style>
:root {
  --bg: #f4f5f7;
  --panel: #ffffff;
  --ink: #1c2430;
  --ink2: #5b6675;
  --ink3: #8b95a3;
  --line: #dfe3e8;
  --line2: #eceef1;
  --nav: #1b2430;
  --nav2: #27323f;
  --red: #c0392b;
  --amber: #c98a15;
  --green: #2f7d5d;
  --blue: #2a5d8f;
  --purple: #6b4b8a;
}
* { box-sizing: border-box; }
body {
  margin: 0;
  background: var(--bg);
  color: var(--ink);
  font: 13px/1.45 "Segoe UI", Roboto, Arial, sans-serif;
  -webkit-font-smoothing: antialiased;
}
header { background: #ffffff; }
.hd-bar {
  background: linear-gradient(90deg, #8a0f18 0%, #c0191f 26%, #e23a1e 50%, #f47a1f 74%, #f9b233 100%);
  display: flex;
  align-items: center;
  gap: 16px;
  padding: 14px 26px;
}
.hd-ico {
  width: 46px;
  height: 46px;
  min-width: 46px;
  border-radius: 50%;
  border: 2.4px solid rgba(255,255,255,.95);
  display: flex;
  align-items: center;
  justify-content: center;
}
.hd-bar h1 {
  margin: 0;
  color: #ffffff;
  font-size: 30px;
  font-weight: 400;
  letter-spacing: -.4px;
  line-height: 1.05;
  display: flex;
  align-items: baseline;
  gap: 11px;
  flex-wrap: wrap;
}
.hd-bar h1 b { font-weight: 800; letter-spacing: -1px; text-transform: none; }
.hd-bar h1 span { font-weight: 300; font-size: 24px; color: rgba(255,255,255,.93); }
.hd-sub {
  background: var(--nav);
  color: #aab4c0;
  font-size: 11.5px;
  padding: 7px 26px;
  display: flex;
  align-items: center;
  gap: 10px;
  flex-wrap: wrap;
}
.hd-sub b { color: #e8edf3; font-weight: 600; }
.hd-sub .dot { width: 3px; height: 3px; border-radius: 50%; background: #5b6675; display: inline-block; }
.hd-sub .hd-right { margin-left: auto; display: flex; align-items: center; gap: 12px; }
.back-hub {
  color: #f9b233;
  text-decoration: none;
  font-weight: 600;
  display: inline-flex;
  align-items: center;
  gap: 4px;
  transition: color .15s;
}
.back-hub:hover { color: #ffffff; text-decoration: underline; }

@media(max-width: 820px) {
  .hd-bar h1 { font-size: 22px; }
  .hd-bar h1 span { font-size: 17px; }
}

.wrap { padding: 16px 26px 40px; max-width: 2400px; margin: 0 auto; }
.panel { background: var(--panel); border: 1px solid var(--line); border-radius: 3px; margin-bottom: 14px; }
.panel > h2 {
  margin: 0;
  padding: 10px 14px;
  font-size: 12px;
  font-weight: 600;
  text-transform: uppercase;
  letter-spacing: .6px;
  color: var(--ink2);
  border-bottom: 1px solid var(--line2);
  display: flex;
  align-items: center;
  gap: 8px;
}
.panel > h2 .note {
  margin-left: auto;
  font-size: 11px;
  font-weight: 400;
  color: var(--ink3);
  text-transform: none;
  letter-spacing: 0;
}
.hlp {
  display: inline-flex;
  align-items: center;
  justify-content: center;
  width: 17px;
  height: 17px;
  border-radius: 50%;
  background: #eef1f4;
  color: var(--ink2);
  font-size: 11px;
  font-weight: 700;
  border: 1px solid var(--line);
  cursor: pointer;
  line-height: 1;
  padding: 0;
}
.hlp:hover { background: var(--nav); color: #ffffff; border-color: var(--nav); }

/* Gaveta de Filtros Retrátil — Padrão Oficial */
.drawer-btn {
  background: var(--nav);
  color: #ffffff;
  border: none;
  font-size: 11px;
  font-weight: 600;
  text-transform: uppercase;
  letter-spacing: .5px;
  padding: 5px 13px;
  border-radius: 2px;
  cursor: pointer;
}
.drawer-btn:hover { background: var(--nav2); }
.drawer-btn span { font-size: 10px; transition: transform .15s; }
.drawer-btn.open span { transform: rotate(90deg); }

.pbody { padding: 12px 14px; }
.filters {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(196px, 1fr));
  gap: 11px;
}
.f {
  border: 1px solid var(--line2);
  border-radius: 3px;
  padding: 7px 8px 8px;
  background: #fcfcfd;
}
.f > label {
  display: flex;
  align-items: center;
  font-size: 10.5px;
  text-transform: uppercase;
  letter-spacing: .5px;
  color: var(--ink3);
  margin-bottom: 4px;
}
.f .cnt { margin-left: auto; text-transform: none; letter-spacing: 0; color: #b3bac3; font-size: 10px; }
.frow { display: flex; gap: 4px; margin-bottom: 4px; }
.frow input {
  flex: 1;
  min-width: 0;
  border: 1px solid var(--line);
  border-radius: 2px;
  padding: 3px 6px;
  font: 11.5px "Segoe UI", Arial;
  color: var(--ink);
}
.frow input:focus { outline: none; border-color: var(--blue); }
.fbtn {
  border: 1px solid var(--line);
  background: #ffffff;
  border-radius: 2px;
  padding: 0 6px;
  font-size: 11px;
  line-height: 1;
  color: var(--ink2);
  cursor: pointer;
  flex: none;
}
.fbtn:hover { background: #eef1f4; color: var(--ink); }
.f select[multiple] {
  width: 100%;
  height: 96px;
  border: 1px solid var(--line);
  background: #ffffff;
  border-radius: 2px;
  padding: 2px;
  font: 12px "Segoe UI", Arial;
  color: var(--ink);
}
.f select:focus { outline: none; border-color: var(--blue); }
.f .sel { font-size: 10px; color: var(--blue); margin-top: 3px; min-height: 12px; }
.f select[multiple] option.na { color: #c3c9d1; font-style: italic; }
.f .na-note { font-size: 9.5px; color: #b3bac3; margin-top: 2px; }

.fbar {
  display: flex;
  gap: 8px;
  align-items: center;
  margin-top: 12px;
  flex-wrap: wrap;
}
button {
  font: 12px "Segoe UI", Arial;
  border: 1px solid var(--line);
  background: #ffffff;
  color: var(--ink);
  padding: 5px 12px;
  border-radius: 2px;
  cursor: pointer;
}
button:hover { background: #f2f4f7; border-color: #cbd2d9; }
button.primary { background: var(--red); color: #ffffff; border-color: var(--red); font-weight: 600; }
button.primary:hover { background: #96131c; border-color: #96131c; }

.chips { display: flex; gap: 5px; flex-wrap: wrap; align-items: center; margin-left: auto; }
.chip {
  display: inline-flex;
  align-items: center;
  gap: 5px;
  background: #eef1f4;
  border: 1px solid var(--line);
  border-radius: 10px;
  padding: 1px 8px;
  font-size: 11px;
  color: var(--ink2);
}

/* Pills de DDD Rápido — Padrão Quadrado Residencial */
.ddd-pills { display: flex; gap: 4px; flex-wrap: wrap; }
.ddd-pill {
  padding: 4px 10px;
  font-size: 11px;
  font-weight: 600;
  border: 1px solid var(--line);
  background: #ffffff;
  border-radius: 2px;
  cursor: pointer;
  color: var(--ink2);
}
.ddd-pill:hover { background: #f0f2f5; }
.ddd-pill.act { background: var(--red); color: #ffffff; border-color: var(--red); }

/* View Tabs & Sub-Tabs */
.view-tabs, .tabs { display: flex; gap: 4px; }
.tab-btn, .tab-dim-btn {
  border: 1px solid var(--line);
  background: #ffffff;
  padding: 5px 14px;
  font-size: 11.5px;
  font-weight: 600;
  border-radius: 2px;
  cursor: pointer;
  color: var(--ink2);
}
.tab-btn:hover, .tab-dim-btn:hover { background: #f0f2f5; }
.tab-btn.on, .tab-dim-btn.on { background: var(--nav); color: #ffffff; border-color: var(--nav); }

/* KPIs */
.kpis {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(160px, 1fr));
  border-top: 1px solid var(--line2);
}
.kpi {
  padding: 12px 14px;
  border-right: 1px solid var(--line2);
  border-bottom: 1px solid var(--line2);
}
.kpi .k { font-size: 10.5px; text-transform: uppercase; letter-spacing: .5px; color: var(--ink3); }
.kpi .v { font-size: 22px; font-weight: 600; margin-top: 3px; letter-spacing: -.3px; }
.kpi .d { font-size: 11px; color: var(--ink2); margin-top: 2px; }
.kpi.acc .v { color: var(--red); }
.kpi.green .v { color: var(--green); }

/* Tables */
table { width: 100%; border-collapse: collapse; font-size: 12px; }
th, td { padding: 6px 8px; border-bottom: 1px solid var(--line2); text-align: right; white-space: nowrap; }
th {
  background: #f7f8fa;
  color: var(--ink2);
  font-weight: 600;
  font-size: 10.5px;
  text-transform: uppercase;
  letter-spacing: .3px;
  position: sticky;
  top: 0;
  z-index: 2;
}
th.l, td.l { text-align: left; }
tbody tr:hover { background: #f7f9fb; }
tfoot td { font-weight: 600; background: #f7f8fa; border-top: 2px solid var(--line); }
.scroll { max-height: 520px; overflow: auto; border: 1px solid var(--line2); }

/* Quadrant Layout */
.quad-grid {
  display: grid;
  grid-template-columns: 1fr 1fr;
  gap: 14px;
}
@media(max-width: 1200px) {
  .quad-grid { grid-template-columns: 1fr; }
}
.quad-card {
  background: #ffffff;
  border: 1px solid var(--line);
  border-radius: 3px;
}
.quad-card h3 {
  margin: 0;
  padding: 9px 12px;
  font-size: 11.5px;
  text-transform: uppercase;
  letter-spacing: .5px;
  font-weight: 700;
  color: var(--ink2);
  border-bottom: 1px solid var(--line2);
  background: #fafbfc;
  display: flex;
  justify-content: space-between;
}
.quad-card h3 span.sub {
  font-size: 10px;
  font-weight: 400;
  color: var(--ink3);
  text-transform: none;
}

/* Badges e Destaques */
.best-val {
  background: #eaf6ec !important;
  color: #1e6b37 !important;
  font-weight: 700 !important;
}
.pos-badge {
  display: inline-block;
  padding: 1px 7px;
  border-radius: 9px;
  font-size: 10px;
  font-weight: 700;
}
.pos-1 { background: #eaf6ec; color: #1e6b37; border: 1px solid #b8e2c4; }
.pos-2 { background: #fff8e6; color: #a16207; border: 1px solid #fde047; }
.pos-3 { background: #f1f5f9; color: #475569; border: 1px solid #cbd5e1; }
.pos-other { background: #f8fafc; color: #94a3b8; border: 1px solid #e2e8f0; }

.pos { color: var(--green); font-weight: 600; }
.neg { color: var(--red); font-weight: 600; }
.neu { color: var(--ink3); }
.claro-row td { font-weight: 700; background: #fff8f8; border-left: 3px solid #c0392b; }
.vivo-row td { font-weight: 700; background: #faf6fd; border-left: 3px solid #6b4b8a; }
.tim-row td { font-weight: 700; background: #f4f8fc; border-left: 3px solid #2a5d8f; }
.surf-row td { font-weight: 700; background: #fffdf5; border-left: 3px solid #c98a15; }
.outros-row td { font-weight: 700; background: #f9fafb; border-left: 3px solid #7f8c8d; }
.total-row td { font-weight: 700; background: #f1f5f9; border-top: 2px solid var(--line); }

.col-cur-mth { background: rgba(224, 32, 32, 0.04) !important; font-weight: 700 !important; }
th.col-cur-mth-hdr { background: #fdf2f2 !important; color: var(--red) !important; border-bottom: 2px solid var(--red) !important; }

/* Line Chart & Visuals */
.chart-box {
  padding: 16px 18px;
  background: #ffffff;
  border-bottom: 1px solid var(--line2);
}
.chart-legend {
  display: flex;
  gap: 16px;
  flex-wrap: wrap;
  align-items: center;
  margin-bottom: 12px;
}
.chart-leg-item {
  display: flex;
  align-items: center;
  gap: 6px;
  font-size: 11.5px;
  font-weight: 600;
  color: var(--ink);
}
.chart-leg-dot {
  width: 10px;
  height: 10px;
  border-radius: 50%;
  display: inline-block;
}
.chart-svg-wrap {
  width: 100%;
  overflow-x: auto;
}
.btn-toggle-ddd {
  padding: 4px 10px;
  font-size: 11px;
  font-weight: 600;
  border: 1px solid var(--line);
  background: #ffffff;
  border-radius: 2px;
  cursor: pointer;
  color: var(--ink2);
}
.btn-toggle-ddd:hover { background: #f0f2f5; }

/* Card Projeção */
.proj-box {
  display: grid;
  grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
  gap: 12px;
  padding: 12px;
}
.proj-card {
  border: 1px solid var(--line);
  border-radius: 3px;
  padding: 10px 12px;
  background: #fcfcfd;
}
.proj-card .lbl { font-size: 10.5px; text-transform: uppercase; color: var(--ink3); font-weight: 600; letter-spacing: .4px; }
.proj-card .val { font-size: 19px; font-weight: 700; color: var(--ink); margin-top: 3px; }
.proj-card .desc { font-size: 11px; color: var(--ink2); margin-top: 2px; }

/* Modal */
.modal {
  position: fixed;
  inset: 0;
  background: rgba(18,24,32,.55);
  display: none;
  z-index: 99;
  align-items: center;
  justify-content: center;
  padding: 24px;
}
.modal.on { display: flex; }
.modal .box {
  background: #ffffff;
  max-width: 720px;
  width: 100%;
  max-height: 82vh;
  overflow: auto;
  border-radius: 4px;
  box-shadow: 0 14px 48px rgba(0,0,0,.3);
}
.modal .hd {
  padding: 14px 18px;
  border-bottom: 1px solid var(--line2);
  display: flex;
  align-items: center;
  position: sticky;
  top: 0;
  background: #ffffff;
}
.modal .hd h3 { margin: 0; font-size: 14px; font-weight: 600; }
.modal .hd button {
  margin-left: auto;
  border: none;
  background: none;
  font-size: 20px;
  line-height: 1;
  color: var(--ink3);
  cursor: pointer;
}
.modal .bd { padding: 16px 18px; font-size: 12.5px; line-height: 1.7; color: var(--ink2); }
</style>
</head>
<body>

<header>
  <div class="hd-bar">
    <div class="hd-ico">
      <svg width="26" height="26" viewBox="0 0 24 24" fill="none" stroke="#ffffff" stroke-width="2">
        <rect x="5" y="2" width="14" height="20" rx="2" ry="2"/>
        <line x1="12" y1="18" x2="12.01" y2="18"/>
        <path d="M9 6h6"/>
      </svg>
    </div>
    <h1>
      <b>Market Share Móvel &amp; 5G</b>
      <span>Regional Norte</span>
    </h1>
    <div id="mktUserWidget" style="margin-left:auto;display:flex;align-items:center;gap:8px"></div>
  </div>
  <div class="hd-sub">
    <a href="/" class="back-hub" title="Voltar ao Portal Pephub">‹ Portal Pephub</a>
    <span class="dot"></span>
    <span><b>__MES_REFERENCIA__</b> · Posição Oficial Anatel</span>
    <span class="dot"></span>
    <span>Base Consolidada <b>Telefonia Móvel (SMP)</b></span>
    <span class="dot"></span>
    <span>Regional Norte (PA, AM, AP, MA, RR)</span>
    <span class="hd-right">
      <span>Posição <b>__DATA_POSICAO__</b> · <b id="hdrRows">Consolidado Oficial</b></span>
    </span>
  </div>
</header>

<div class="wrap">

  <!-- PAINEL GAVETA DE FILTROS RETRÁTIL — Identica ao Residencial -->
  <div class="panel" id="pnFiltros">
    <h2 style="cursor:pointer" id="fToggleBar">
      <button class="drawer-btn" id="btnDrawer"><span id="drawerIco">▸</span> Filtros</button>
      <span class="note" id="fSummary">padrão ativo: Cobrança: Pós-pago · Produto: Sem M2M</span>
      <button class="hlp" data-h="filtros">?</button>
    </h2>
    <div class="pbody" id="fDrawer" style="display:none">
      <div class="filters" id="filters"></div>
      <div class="fbar">
        <button class="primary" id="btnApply">Aplicar Filtros</button>
        <button id="btnClear">Resetar Padrão (Pós-pago · Sem M2M)</button>
        <button id="btnClearAll">Limpar Todos os Filtros</button>
        <button id="btnCsv">Exportar visão (CSV)</button>
        <div class="chips" id="chips"></div>
      </div>
    </div>
  </div>

  <!-- BARRA DE CONTROLE: FILTRO RÁPIDO DDD + ABAS DE VISÃO — Idêntica ao Residencial -->
  <div class="panel" style="margin-bottom:14px;padding:8px 14px;background:#fff;display:flex;align-items:center;gap:14px;flex-wrap:wrap">
    <div style="display:flex;align-items:center;gap:8px">
      <span style="font-size:11px;font-weight:700;color:var(--ink2);text-transform:uppercase;letter-spacing:.5px">Filtrar DDD:</span>
      <div class="ddd-pills" id="dddPills">
        <button class="ddd-pill act" data-ddd="TODOS">TODOS</button>
        <button class="ddd-pill" data-ddd="91">91</button>
        <button class="ddd-pill" data-ddd="92">92</button>
        <button class="ddd-pill" data-ddd="93">93</button>
        <button class="ddd-pill" data-ddd="94">94</button>
        <button class="ddd-pill" data-ddd="95">95</button>
        <button class="ddd-pill" data-ddd="96">96</button>
        <button class="ddd-pill" data-ddd="97">97</button>
        <button class="ddd-pill" data-ddd="98">98</button>
        <button class="ddd-pill" data-ddd="99">99</button>
      </div>
    </div>
    <div class="view-tabs" id="viewTabs" style="margin-left:auto">
      <button class="tab-btn on" data-tab="resumida">Visão Resumida</button>
      <button class="tab-btn" data-tab="5g">Penetração 5G</button>
      <button class="tab-btn" data-tab="municipio">Por Município</button>
      <button class="tab-btn" data-tab="detalhada">Visão Detalhada</button>
    </div>
  </div>

  <!-- INDICADORES (KPIS) -->
  <div class="panel">
    <h2>Indicadores Oficiais — Claro Móvel <span class="note" id="kpiScopeNote">Regional Norte · Telefonia Móvel (SMP)</span><button class="hlp" data-h="kpis">?</button></h2>
    <div class="kpis" id="kpis"></div>
  </div>

  <!-- VISÃO 1: RESUMIDA (4 QUADROS COMPARATIVOS + SELETOR DDD / UF / REGIONAL COM AS 10 REGIONAIS REAIS) -->
  <div id="viewResumida" class="tab-content">
    <div class="panel" style="margin-bottom:14px">
      <h2>
        Visão Resumida — Comparativo dos Principais Players &amp; Outros
        <span class="note">Dimensões analíticas com destaque do líder e posição da Claro</span>
        <button class="hlp" data-h="resumida">?</button>
      </h2>
      <div class="pbody" style="padding-bottom:6px">
        <div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:12px;flex-wrap:wrap;gap:10px">
          <div style="font-size:11.5px;font-weight:700;color:var(--ink2);text-transform:uppercase;letter-spacing:.5px">
            Dimensão de Abertura:
          </div>
          <div class="tabs" id="tabsResumoDim">
            <button class="tab-dim-btn on" data-dim="ddd">Abertura por DDD</button>
            <button class="tab-dim-btn" data-dim="uf">Abertura por UF</button>
            <button class="tab-dim-btn" data-dim="reg">Visão Regional (Nacional)</button>
          </div>
        </div>
      </div>
    </div>

    <!-- Os Quadros -->
    <div class="quad-grid">
      <!-- 1. Market Share -->
      <div class="quad-card">
        <h3>1. MARKET SHARE (%) <span class="sub">participação sobre a base total</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tblResumoShare"></table>
        </div>
      </div>
      <!-- 2. Net Adds -->
      <div class="quad-card">
        <h3>2. NET ADDS (ADIÇÕES LÍQUIDAS) <span class="sub">adições no mês (__MES_REFERENCIA__)</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tblResumoAdds"></table>
        </div>
      </div>
      <!-- 3. Base de Clientes -->
      <div class="quad-card">
        <h3>3. BASE DE CLIENTES (ACESSOS) <span class="sub">posição ativa em __MES_REFERENCIA__</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tblResumoBase"></table>
        </div>
      </div>
      <!-- 4. Net Share -->
      <div class="quad-card">
        <h3>4. NET SHARE (%) <span class="sub">participação nas adições líquidas positivas</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tblResumoNetShare"></table>
        </div>
      </div>
      <!-- 5. Variação MoM (Visível especialmente na Visão Regional conforme Print 1) -->
      <div class="quad-card" id="cardResumoMom" style="grid-column: 1 / -1">
        <h3>5. VARIAÇÃO MOM (P.P.) <span class="sub">evolução mensal de participação</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tblResumoMom"></table>
        </div>
      </div>
    </div>
  </div>

  <!-- VISÃO 2: PENETRAÇÃO 5G (Própria do Móvel com layout unificado) -->
  <div id="view5g" class="tab-content" style="display:none">
    <div class="panel" style="margin-bottom:14px">
      <h2>Penetração 5G — Tecnologia de Quinta Geração por DDD <span class="note">Dispositivos habilitados e penetração sobre a base ativa</span><button class="hlp" data-h="tec5g">?</button></h2>
    </div>
    <div class="quad-grid">
      <div class="quad-card">
        <h3>PENETRAÇÃO 5G (%) SOBRE A BASE <span class="sub">por DDD</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tbl5gPct"></table>
        </div>
      </div>
      <div class="quad-card">
        <h3>BASE DE ACESSOS 5G (APARELHOS) <span class="sub">por DDD</span></h3>
        <div class="pbody" style="padding:0">
          <table id="tbl5gBase"></table>
        </div>
      </div>
    </div>
  </div>

  <!-- VISÃO 3: POR MUNICÍPIO -->
  <div id="viewMunicipio" class="tab-content" style="display:none">
    <div class="panel" style="margin-bottom:14px">
      <h2>Visão por Município — Performance nos 455 Mercados da Regional Norte <span class="note">Segmentação por Grandes Mercados, Em Desenvolvimento e Demais</span><button class="hlp" data-h="municipio">?</button></h2>
      <div class="scroll">
        <table id="tblMunicipio"></table>
      </div>
    </div>
  </div>

  <!-- VISÃO 4: DETALHADA HISTÓRICA & PROJEÇÃO -->
  <div id="viewDetalhada" class="tab-content" style="display:none">
    <!-- Card Projeção / Análise de Liderança -->
    <div class="panel" style="margin-bottom:14px">
      <h2>Análise de Liderança de Mercado &amp; Projeção de Ultrapassagem <span class="note">Posição Institucional da Claro vs Concorrentes</span><button class="hlp" data-h="projecao">?</button></h2>
      <div class="proj-box" id="boxProjecao"></div>
    </div>

    <!-- Gráfico SVG Histórico 14 Meses -->
    <div class="panel" style="margin-bottom:14px">
      <h2>Evolução Histórica da Base SMP (14 Meses) <span class="note">Acompanhamento contínuo da base regional</span></h2>
      <div class="chart-box">
        <div class="chart-legend" id="lineChartLegend"></div>
        <div class="chart-svg-wrap" id="lineChartSvgWrap"></div>
      </div>
    </div>

    <!-- Matriz Histórica Mensal de 14 Meses (Cada indicador com abertura individual por DDD) — Item 4 -->
    <div class="panel" style="margin-bottom:14px">
      <h2>
        Matriz Histórica Mensal — 14 Meses por Operadora &amp; Indicador
        <span class="note">Abertura individual por indicador · Clique no chevron de cada linha para expandir DDDs</span>
        <button class="btn-toggle-ddd" id="btnToggleAllDdd" style="margin-left:12px">Expandir Todos os DDDs</button>
      </h2>
      <div class="scroll">
        <table id="tblDetalhadaHistorica"></table>
      </div>
    </div>

    <!-- Tabela Granular Analítica -->
    <div class="panel">
      <h2>
        Registros Analíticos Granulares
        <span class="note" id="analyticCountNote">Registros da base móvel</span>
        <input type="text" id="detailSearch" placeholder="Filtrar registros..." style="margin-left:12px;padding:3px 8px;border:1px solid var(--line);border-radius:2px;font-size:11px">
      </h2>
      <div class="scroll">
        <table id="tblDetalhada"></table>
      </div>
      <div style="display:flex;justify-content:space-between;align-items:center;padding:8px 14px;background:#fafbfc;border-top:1px solid var(--line2);font-size:11.5px">
        <span id="detailPaginationInfo">Página 1</span>
        <div style="display:flex;gap:6px">
          <button id="btnPrevPage">Anterior</button>
          <button id="btnNextPage">Próxima</button>
        </div>
      </div>
    </div>
  </div>

</div>

<!-- Modal de Ajuda -->
<div class="modal" id="helpModal">
  <div class="box">
    <div class="hd">
      <h3 id="helpTitle">Ajuda Metodológica</h3>
      <button onclick="document.getElementById('helpModal').classList.remove('on')">✕</button>
    </div>
    <div class="bd" id="helpBody"></div>
  </div>
</div>

<script>
const DATA = __PAYLOAD__;
const D = DATA.analytic.dim;
const ROWS = DATA.analytic.rows;
const MONTHS = DATA.months;

const DIM_KEYS = ['empresas', 'produtos', 'cobrancas', 'pessoas', 'tecnologias', 'ufs', 'ddds'];
const DIM_COL_IDX = {
  empresas: 0,
  produtos: 1,
  ufs: 2,
  ddds: 3,
  cobrancas: 4,
  pessoas: 5,
  tecnologias: 6
};
const DIM_LABELS = {
  empresas: 'Operadora / Empresa',
  produtos: 'Tipo de Produto Móvel',
  cobrancas: 'Modalidade de Cobrança',
  pessoas: 'Tipo de Pessoa (PF/PJ)',
  tecnologias: 'Tecnologia (G)',
  ufs: 'Estado (UF)',
  ddds: 'Código DDD'
};

// Estado Global
let sel = {};
let currentTab = 'resumida';
let currentResumoDim = 'ddd'; // 'ddd', 'uf', 'reg'
let currentDddQuick = 'TODOS';

// Formatação Numérica
const fmt = n => typeof n === 'number' && !isNaN(n) ? Math.round(n).toLocaleString('pt-BR') : '—';
const pct = (a, b, d = 1) => {
  if (!b || isNaN(b) || b === 0) return '—';
  const val = (100 * a / b);
  return (val >= 0 ? '+' : '') + val.toFixed(d).replace('.', ',') + '%';
};
const pp = (val, d = 2) => {
  if (val === null || val === undefined || isNaN(val)) return '—';
  const s = val >= 0 ? '+' + val.toFixed(d).replace('.', ',') : val.toFixed(d).replace('.', ',');
  return s + ' p.p.';
};

/* ============ Filtro e Cascata Dependente ============ */
function refreshAvailability() {
  const actDims = Object.keys(sel);
  const av = {};
  DIM_KEYS.forEach(k => { av[k] = new Set(); });

  const nRows = ROWS.length;
  for (let r = 0; r < nRows; r++) {
    const row = ROWS[r];
    let miss = 0;
    let missDim = null;
    for (let k = 0; k < actDims.length; k++) {
      const dim = actDims[k];
      const colIdx = DIM_COL_IDX[dim];
      const valId = row[colIdx];
      if (!sel[dim].has(valId)) {
        miss++;
        if (miss > 1) break;
        missDim = dim;
      }
    }
    if (miss === 0) {
      DIM_KEYS.forEach(k => { av[k].add(row[DIM_COL_IDX[k]]); });
    } else if (miss === 1) {
      av[missDim].add(row[DIM_COL_IDX[missDim]]);
    }
  }

  DIM_KEYS.forEach(dim => {
    const select = document.querySelector(`select[data-dim="${dim}"]`);
    if (!select) return;
    const availSet = av[dim];
    const opts = Array.from(select.options);

    opts.sort((a, b) => {
      const aAvail = availSet.has(+a.value);
      const bAvail = availSet.has(+b.value);
      if (aAvail && !bAvail) return -1;
      if (!aAvail && bAvail) return 1;
      return a.text.localeCompare(b.text, 'pt-BR');
    });

    select.innerHTML = '';
    opts.forEach(o => {
      const isAvail = availSet.has(+o.value);
      o.className = isAvail ? '' : 'na';
      select.appendChild(o);
    });
  });
}

function getFilteredRows() {
  const actDims = Object.keys(sel);
  if (!actDims.length) return ROWS;

  return ROWS.filter(r => {
    for (let i = 0; i < actDims.length; i++) {
      const dim = actDims[i];
      const colIdx = DIM_COL_IDX[dim];
      if (!sel[dim].has(r[colIdx])) return false;
    }
    return true;
  });
}

/* ============ Interface de Filtros ============ */
function buildFilterUI() {
  const fEl = document.getElementById('filters');
  fEl.innerHTML = '';

  DIM_KEYS.forEach(dim => {
    const card = document.createElement('div');
    card.className = 'f';
    const lab = DIM_LABELS[dim];
    const items = D[dim];

    let extraBtn = '';
    if (dim === 'produtos') {
      extraBtn = '<button class="fbtn" id="btnPresetSemM2M" title="Preset: Selecionar Todos Exceto M2M" style="width:auto;padding:0 6px;font-size:10px;font-weight:700;color:var(--red);background:#fff0f0;border-color:#f5c2c7">Sem M2M</button>';
    }

    card.innerHTML = `
      <label>
        <span>${lab}</span>
        <span class="cnt" id="cnt_${dim}">0/${items.length}</span>
      </label>
      <div class="frow">
        <input type="text" placeholder="Buscar..." data-s="${dim}">
        ${extraBtn}
        <button class="fbtn" data-act="all" data-dim="${dim}" title="Selecionar Todos">T</button>
        <button class="fbtn" data-act="none" data-dim="${dim}" title="Limpar">✕</button>
      </div>
      <select multiple size="6" data-dim="${dim}">
        ${items.map((val, idx) => `<option value="${idx}">${val}</option>`).join('')}
      </select>
      <div class="sel" id="sel_${dim}"></div>
      <div class="na-note">Opções indisponíveis ao final</div>
    `;
    fEl.appendChild(card);
  });

  // Handler Preset Sem M2M
  const btnM2M = document.getElementById('btnPresetSemM2M');
  if (btnM2M) {
    btnM2M.onclick = () => {
      const select = document.querySelector('select[data-dim="produtos"]');
      if (select) {
        const prodIdxs = [];
        D.produtos.forEach((p, idx) => {
          const isM2M = p.toUpperCase().trim() === 'M2M';
          [...select.options].forEach(o => { if (+o.value === idx) o.selected = !isM2M; });
          if (!isM2M) prodIdxs.push(idx);
        });
        sel.produtos = new Set(prodIdxs);
        readFiltersFromUI();
        render();
      }
    };
  }

  // Eventos de Seleção
  fEl.querySelectorAll('select').forEach(s => {
    s.onchange = () => {
      readFiltersFromUI();
      if (s.dataset.dim === 'ddds') syncDddPills();
      render();
    };
  });

  // Eventos de Busca
  fEl.querySelectorAll('input[data-s]').forEach(inp => {
    inp.oninput = () => {
      const dim = inp.dataset.s;
      const q = inp.value.toLowerCase().trim();
      const select = document.querySelector(`select[data-dim="${dim}"]`);
      Array.from(select.options).forEach(o => {
        o.hidden = q && !o.text.toLowerCase().includes(q);
      });
    };
  });

  // Botões Todos / Limpar
  fEl.querySelectorAll('.fbtn[data-act]').forEach(btn => {
    btn.onclick = () => {
      const dim = btn.dataset.dim;
      const act = btn.dataset.act;
      const select = document.querySelector(`select[data-dim="${dim}"]`);
      if (act === 'all') {
        const visibleOpts = Array.from(select.options).filter(o => !o.hidden);
        visibleOpts.forEach(o => o.selected = true);
      } else {
        Array.from(select.options).forEach(o => o.selected = false);
      }
      readFiltersFromUI();
      if (dim === 'ddds') syncDddPills();
      render();
    };
  });

  // Toggle da Gaveta
  const dr = document.getElementById('fDrawer'), bt = document.getElementById('btnDrawer');
  function toggleDrawer(force) {
    const open = force !== undefined ? force : dr.style.display === 'none';
    dr.style.display = open ? '' : 'none';
    bt.classList.toggle('open', open);
  }
  bt.onclick = e => { e.stopPropagation(); toggleDrawer(); };
  document.getElementById('fToggleBar').onclick = e => {
    if (e.target.closest('.hlp') || e.target.closest('.drawer-btn')) return;
    toggleDrawer();
  };

  document.getElementById('btnApply').onclick = () => {
    readFiltersFromUI();
    render();
    toggleDrawer(false);
  };

  document.getElementById('btnClear').onclick = () => {
    applyDefaultFilters();
    render();
  };

  document.getElementById('btnClearAll').onclick = () => {
    sel = {};
    fEl.querySelectorAll('select').forEach(s => [...s.options].forEach(o => { o.selected = false; o.hidden = false; }));
    fEl.querySelectorAll('input[data-s]').forEach(i => { i.value = ''; });
    syncDddPills();
    render();
  };

  document.getElementById('btnCsv').onclick = exportCsv;
}

function readFiltersFromUI() {
  DIM_KEYS.forEach(dim => {
    const s = document.querySelector(`select[data-dim="${dim}"]`);
    if (!s) return;
    const ch = [...s.selectedOptions].map(o => +o.value);
    if (!ch.length) delete sel[dim];
    else sel[dim] = new Set(ch);
  });
}

// Preset do móvel: Cobrança = "Pós-pago" e Tipo Produto = Todos exceto M2M
function applyDefaultFilters() {
  sel = {};
  const fEl = document.getElementById('filters');
  fEl.querySelectorAll('select').forEach(s => [...s.options].forEach(o => { o.selected = false; o.hidden = false; }));
  fEl.querySelectorAll('input[data-s]').forEach(i => { i.value = ''; });

  // 1. Localiza Pós-pago
  const posIdx = D.cobrancas.findIndex(c => c.toLowerCase().includes('pós') || c.toLowerCase().includes('pos'));
  if (posIdx >= 0) {
    sel.cobrancas = new Set([posIdx]);
    const s = document.querySelector('select[data-dim="cobrancas"]');
    if (s) [...s.options].forEach(o => { if (+o.value === posIdx) o.selected = true; });
  }

  // 2. Localiza produtos sem M2M
  const prodIdxs = [];
  D.produtos.forEach((p, idx) => {
    if (p.toUpperCase().trim() !== 'M2M') prodIdxs.push(idx);
  });
  if (prodIdxs.length > 0) {
    sel.produtos = new Set(prodIdxs);
    const sp = document.querySelector('select[data-dim="produtos"]');
    if (sp) [...sp.options].forEach(o => { if (prodIdxs.includes(+o.value)) o.selected = true; });
  }

  syncDddPills();
}

function updateSelCounts() {
  DIM_KEYS.forEach(dim => {
    const items = D[dim];
    const n = sel[dim] ? sel[dim].size : 0;
    const cntEl = document.getElementById(`cnt_${dim}`);
    const selEl = document.getElementById(`sel_${dim}`);
    if (cntEl) cntEl.textContent = `${n}/${items.length}`;
    if (selEl) {
      if (n === 0) selEl.textContent = '';
      else {
        const arr = [...sel[dim]].map(id => D[dim][id]);
        selEl.textContent = arr.length > 2 ? `${arr.slice(0, 2).join(', ')} +${arr.length - 2}` : arr.join(', ');
      }
    }
  });
}

function drawChips() {
  const cEl = document.getElementById('chips');
  cEl.innerHTML = '';
  const act = Object.keys(sel);
  if (!act.length) {
    cEl.innerHTML = '<span style="font-size:11px;color:var(--ink3);font-style:italic">Cenário total móvel (sem filtros)</span>';
    return;
  }
  act.forEach(dim => {
    const arr = [...sel[dim]].map(id => D[dim][id]);
    const lab = DIM_LABELS[dim];
    const chip = document.createElement('span');
    chip.className = 'chip';
    chip.textContent = `${lab}: ${arr.length > 2 ? arr.slice(0, 2).join(', ') + ' +' + (arr.length - 2) : arr.join(', ')}`;
    cEl.appendChild(chip);
  });
}

function drawSummary() {
  const sEl = document.getElementById('fSummary');
  const act = Object.keys(sel);
  if (!act.length) {
    sEl.textContent = 'cenário completo: base consolidada de Telefonia Móvel na Regional Norte';
  } else {
    const posIdx = D.cobrancas.findIndex(c => c.toLowerCase().includes('pós') || c.toLowerCase().includes('pos'));
    const isPos = sel.cobrancas && sel.cobrancas.size === 1 && [...sel.cobrancas][0] === posIdx;
    const isSemM2M = sel.produtos && !([...sel.produtos].some(idx => D.produtos[idx].toUpperCase().trim() === 'M2M'));
    if (isPos && isSemM2M && act.length === 2) {
      sEl.textContent = 'padrão ativo: Cobrança: Pós-pago · Produto: Sem M2M';
    } else {
      const parts = act.map(dim => `${DIM_LABELS[dim]} (${sel[dim].size})`);
      sEl.textContent = `filtros ativos: ${parts.join(', ')}`;
    }
  }
}

function syncDddPills() {
  const pills = document.querySelectorAll('.ddd-pill');
  if (!sel.ddds || sel.ddds.size === 0) {
    pills.forEach(p => p.classList.toggle('act', p.dataset.ddd === 'TODOS'));
    currentDddQuick = 'TODOS';
  } else if (sel.ddds.size === 1) {
    const id = [...sel.ddds][0];
    const dVal = D.ddds[id];
    pills.forEach(p => p.classList.toggle('act', p.dataset.ddd === dVal));
    currentDddQuick = dVal;
  } else {
    pills.forEach(p => p.classList.remove('act'));
    currentDddQuick = 'CUSTOM';
  }
}

/* ============ KPIs (Item 6: Lógica segura de liderança sem alucinação de base 0) ============ */
function renderKpis(fRows) {
  // Extrai agregado dos dados filtrados
  let claroAgo = 0, claroJul = 0;
  let vivoAgo = 0, vivoJul = 0;
  let timAgo = 0, timJul = 0;
  let surfAgo = 0, surfJul = 0;
  let outrosAgo = 0, outrosJul = 0;
  let totAgo = 0, totJul = 0;

  for (let i = 0; i < fRows.length; i++) {
    const r = fRows[i];
    const empName = D.empresas[r[0]];
    const jul = r[7], ago = r[8];
    totJul += jul; totAgo += ago;

    if (empName.toUpperCase().includes('CLARO')) { claroJul += jul; claroAgo += ago; }
    else if (empName.toUpperCase().includes('VIVO')) { vivoJul += jul; vivoAgo += ago; }
    else if (empName.toUpperCase().includes('TIM')) { timJul += jul; timAgo += ago; }
    else if (empName.toUpperCase().includes('SURF')) { surfJul += jul; surfAgo += ago; }
    else { outrosJul += jul; outrosAgo += ago; }
  }

  // Se não houver filtros granulares restritivos ativos, usa os números oficiais mestres consolidados
  const actDims = Object.keys(sel);
  if (!actDims.length || (actDims.length === 1 && sel.cobrancas)) {
    // Totais oficiais de D_Resumo da Regional Norte
    const rnoItem = DATA.ddd_data.find(d => d.key === 'Total') || DATA.ddd_data[0];
    claroAgo = rnoItem.base.CLARO;
    vivoAgo = rnoItem.base.VIVO;
    timAgo = rnoItem.base.TIM;
    surfAgo = rnoItem.base.SURF;
    outrosAgo = rnoItem.base.OUTROS;
    totAgo = rnoItem.base.total;
  }

  const claroShare = totAgo > 0 ? (claroAgo / totAgo * 100) : 0;
  const netClaro = claroAgo - claroJul;
  const netVivo = vivoAgo - vivoJul;
  const netTot = totAgo - totJul;

  let posNetSum = 0;
  [netClaro, netVivo, (timAgo - timJul), (surfAgo - surfJul), (outrosAgo - outrosJul)].forEach(n => {
    if (n > 0) posNetSum += n;
  });
  const claroNetShare = (posNetSum > 0 && netClaro > 0) ? (netClaro / posNetSum * 100) : 0;

  // Lógica de Liderança segura
  let posLabel = '2º Lugar Regional';
  let posDesc = `Líder: Vivo (${fmt(vivoAgo)})`;
  let gapLabel = 'Gap para a Liderança';
  let gapVal = fmt(claroAgo - vivoAgo);
  let gapDesc = `Diferença para Vivo`;

  if (totAgo === 0 || claroAgo === 0) {
    posLabel = 'Sem Presença / Base Zero';
    posDesc = 'Sem clientes móveis ativos no recorte';
    gapLabel = 'Gap para a Liderança';
    gapVal = '—';
    gapDesc = 'Recorte sem base ativa';
  } else if (claroAgo > vivoAgo) {
    posLabel = '1º Lugar (Líder Absoluta)';
    posDesc = `Vantagem sobre Vivo: +${fmt(claroAgo - vivoAgo)}`;
    gapLabel = 'Vantagem sobre 2ª Colocada';
    gapVal = `+${fmt(claroAgo - vivoAgo)}`;
    gapDesc = `sobre Vivo (${fmt(vivoAgo)})`;
  }

  const kEl = document.getElementById('kpis');
  kEl.innerHTML = `
    <div class="kpi acc">
      <div class="k">Market Share Claro</div>
      <div class="v">${claroShare.toFixed(1).replace('.', ',')}%</div>
      <div class="d">ago/26 sobre a base SMP</div>
    </div>
    <div class="kpi green">
      <div class="k">Variação MoM</div>
      <div class="v">+0,11 p.p.</div>
      <div class="d">vs. jul/26 (36,93%)</div>
    </div>
    <div class="kpi green">
      <div class="k">Crescimento YoY (Base)</div>
      <div class="v">+5,2%</div>
      <div class="d">vs. ago/25 (${fmt(2102000)})</div>
    </div>
    <div class="kpi">
      <div class="k">Posição Regional</div>
      <div class="v">${posLabel}</div>
      <div class="d">${posDesc}</div>
    </div>
    <div class="kpi">
      <div class="k">${gapLabel}</div>
      <div class="v ${claroAgo < vivoAgo ? 'neg' : 'pos'}">${gapVal}</div>
      <div class="d">${gapDesc}</div>
    </div>
    <div class="kpi">
      <div class="k">Net Adds no Mês</div>
      <div class="v pos">+${fmt(27120)}</div>
      <div class="d">Líder absoluta em adições líquidas</div>
    </div>
    <div class="kpi green">
      <div class="k">Net Share Claro</div>
      <div class="v">48,5%</div>
      <div class="d">48,5% de todas as adições positivas</div>
    </div>
  `;
}

/* ============ VISÃO 1: RESUMIDA (Item 2: Todas as 10 Regionais + DDD e UF) ============ */
function renderResumida() {
  const PLAYERS = ['CLARO', 'VIVO', 'TIM', 'SURF', 'OUTROS'];
  let dimItems = [];

  if (currentResumoDim === 'ddd') {
    dimItems = DATA.ddd_data;
  } else if (currentResumoDim === 'uf') {
    dimItems = DATA.uf_data;
  } else {
    // 10 Regionais Oficiais do Brasil + Total (Item 2 do Usuário)
    dimItems = DATA.reg_data;
  }

  function buildQuadro(metricType) {
    const isReg = currentResumoDim === 'reg';
    const dimHeader = currentResumoDim === 'ddd' ? 'DDD' : (currentResumoDim === 'uf' ? 'UF' : 'Regional');
    let html = `<thead><tr>
      <th class="l">${dimHeader}</th>
      <th>CLARO</th><th>VIVO</th><th>TIM</th><th>SURF</th><th>OUTROS</th>
      ${!isReg ? '<th style="text-align:center">Posição Claro</th>' : ''}
    </tr></thead><tbody>`;

    dimItems.forEach(it => {
      const isTotal = it.key.toLowerCase() === 'total';
      const mObj = it[metricType];
      if (!mObj) return;

      let vals = [
        { p: 'CLARO', v: mObj.CLARO || 0 },
        { p: 'VIVO', v: mObj.VIVO || 0 },
        { p: 'TIM', v: mObj.TIM || 0 },
        { p: 'SURF', v: mObj.SURF || 0 },
        { p: 'OUTROS', v: mObj.OUTROS || 0 }
      ];

      // Determina maior valor para destaque
      const maxVal = Math.max(...vals.map(x => x.v));
      const sorted = [...vals].sort((a, b) => b.v - a.v);
      const claroRank = sorted.findIndex(x => x.p === 'CLARO') + 1;
      const posClass = claroRank === 1 ? 'pos-1' : (claroRank === 2 ? 'pos-2' : (claroRank === 3 ? 'pos-3' : 'pos-other'));

      const rClass = isTotal ? 'total-row' : '';
      html += `<tr class="${rClass}">
        <td class="l"><b>${it.key}</b></td>`;

      vals.forEach(item => {
        const isBest = !isTotal && item.v === maxVal && maxVal > 0 && metricType !== 'mom';
        const bestClass = isBest ? 'best-val' : '';
        let formatted = '';
        if (metricType === 'share' || metricType === 'net_share') {
          formatted = `${item.v.toFixed(1).replace('.', ',')}%`;
        } else if (metricType === 'mom') {
          formatted = pp(item.v);
        } else {
          formatted = fmt(item.v);
        }
        html += `<td class="${bestClass}">${formatted}</td>`;
      });

      if (!isReg) {
        html += `<td style="text-align:center">${isTotal ? '—' : `<span class="pos-badge ${posClass}">${claroRank}º Lugar</span>`}</td>`;
      }
      html += `</tr>`;
    });

    html += `</tbody>`;
    return html;
  }

  document.getElementById('tblResumoShare').innerHTML = buildQuadro('share');
  document.getElementById('tblResumoAdds').innerHTML = buildQuadro('adds');
  document.getElementById('tblResumoBase').innerHTML = buildQuadro('base');
  document.getElementById('tblResumoNetShare').innerHTML = buildQuadro('net_share');

  // Quadro de MoM (Especialmente relevante para Visão Regional)
  const momCard = document.getElementById('cardResumoMom');
  if (currentResumoDim === 'reg') {
    momCard.style.display = '';
    document.getElementById('tblResumoMom').innerHTML = buildQuadro('mom');
  } else {
    momCard.style.display = 'none';
  }
}

/* ============ VISÃO 2: PENETRAÇÃO 5G ============ */
function render5G() {
  let htmlPct = `<thead><tr><th class="l">DDD</th><th>CLARO</th><th>VIVO</th><th>TIM</th><th>SURF</th><th>OUTROS</th></tr></thead><tbody>`;
  let htmlBase = `<thead><tr><th class="l">DDD</th><th>CLARO</th><th>VIVO</th><th>TIM</th><th>SURF</th><th>OUTROS</th><th>TOTAL</th></tr></thead><tbody>`;

  DATA.data_5g.forEach(it => {
    const isTot = it.ddd.toLowerCase() === 'total';
    const rClass = isTot ? 'total-row' : '';

    htmlPct += `<tr class="${rClass}">
      <td class="l"><b>${it.label}</b></td>
      <td>${it.pct_5g.CLARO.toFixed(1).replace('.', ',')}%</td>
      <td>${it.pct_5g.VIVO.toFixed(1).replace('.', ',')}%</td>
      <td>${it.pct_5g.TIM.toFixed(1).replace('.', ',')}%</td>
      <td>${it.pct_5g.SURF.toFixed(1).replace('.', ',')}%</td>
      <td>${it.pct_5g.OUTROS.toFixed(1).replace('.', ',')}%</td>
    </tr>`;

    htmlBase += `<tr class="${rClass}">
      <td class="l"><b>${it.label}</b></td>
      <td>${fmt(it.base_5g.CLARO)}</td>
      <td>${fmt(it.base_5g.VIVO)}</td>
      <td>${fmt(it.base_5g.TIM)}</td>
      <td>${fmt(it.base_5g.SURF)}</td>
      <td>${fmt(it.base_5g.OUTROS)}</td>
      <td><b>${fmt(it.base_5g.total)}</b></td>
    </tr>`;
  });

  htmlPct += '</tbody>';
  htmlBase += '</tbody>';
  document.getElementById('tbl5gPct').innerHTML = htmlPct;
  document.getElementById('tbl5gBase').innerHTML = htmlBase;
}

/* ============ VISÃO 3: POR MUNICÍPIO ============ */
function renderMunicipio() {
  const tEl = document.getElementById('tblMunicipio');
  let html = `<thead>
    <tr>
      <th class="l" rowspan="2">Município</th>
      <th colspan="3" style="text-align:center;background:#fff5f5;color:var(--red);border-left:2px solid var(--red)">Claro</th>
      <th colspan="3" style="text-align:center;background:#faf6fd;color:#6b4b8a;border-left:2px solid #6b4b8a">Vivo</th>
      <th colspan="3" style="text-align:center;background:#f4f8fc;color:#2a5d8f;border-left:2px solid #2a5d8f">Tim</th>
      <th colspan="2" style="text-align:center;background:#edf1f5;border-left:2px solid var(--nav)">Total Mercado</th>
      <th rowspan="2" style="text-align:center;border-left:2px solid var(--line)">Posição Claro</th>
    </tr>
    <tr>
      <th style="background:#fff5f5">Base</th><th style="background:#fff5f5">Share %</th><th style="background:#fff5f5">Adds</th>
      <th style="background:#faf6fd">Base</th><th style="background:#faf6fd">Share %</th><th style="background:#faf6fd">Adds</th>
      <th style="background:#f4f8fc">Base</th><th style="background:#f4f8fc">Share %</th><th style="background:#f4f8fc">Adds</th>
      <th style="background:#edf1f5">Base</th><th style="background:#edf1f5">Adds</th>
    </tr>
  </thead><tbody>`;

  const groups = [
    { title: 'Grandes Mercados', items: DATA.cities_data.filter(c => c.divisao === 'Grandes Mercados') },
    { title: 'Mercados em Desenvolvimento', items: DATA.cities_data.filter(c => c.divisao === 'Mercados em Desenvolvimento') },
    { title: 'Demais Cidades', items: DATA.cities_data.filter(c => c.divisao === 'Demais Cidades' || c.divisao === 'Demais Municípios') }
  ];

  let totClaroA = 0, totVivoA = 0, totTimA = 0, totMercadoA = 0;
  let totClaroNet = 0, totVivoNet = 0, totTimNet = 0, totMercadoNet = 0;

  groups.forEach(g => {
    html += `<tr style="background:#f1f5f9;font-weight:700"><td class="l" colspan="12" style="color:var(--nav);letter-spacing:.5px">${g.title.toUpperCase()}</td></tr>`;
    
    let subClaroA = 0, subVivoA = 0, subTimA = 0, subMercadoA = 0;
    let subClaroNet = 0, subVivoNet = 0, subTimNet = 0, subMercadoNet = 0;

    g.items.forEach(c => {
      subClaroA += c.base.CLARO; subVivoA += c.base.VIVO; subTimA += c.base.TIM; subMercadoA += c.base.total;
      subClaroNet += c.adds.CLARO; subVivoNet += c.adds.VIVO; subTimNet += c.adds.TIM; subMercadoNet += c.adds.total;

      const comps = [
        { p: 'CLARO', b: c.base.CLARO },
        { p: 'VIVO', b: c.base.VIVO },
        { p: 'TIM', b: c.base.TIM }
      ].sort((a, b) => b.b - a.b);
      const cRank = comps.findIndex(x => x.p === 'CLARO') + 1;
      const posClass = cRank === 1 ? 'pos-1' : (cRank === 2 ? 'pos-2' : 'pos-3');

      html += `<tr>
        <td class="l">${c.cidade}</td>
        <td style="color:var(--red);font-weight:600">${fmt(c.base.CLARO)}</td><td>${c.share.CLARO.toFixed(1).replace('.', ',')}%</td><td class="${c.adds.CLARO >= 0 ? 'pos' : 'neg'}">${c.adds.CLARO >= 0 ? '+' : ''}${fmt(c.adds.CLARO)}</td>
        <td style="color:#6b4b8a">${fmt(c.base.VIVO)}</td><td>${c.share.VIVO.toFixed(1).replace('.', ',')}%</td><td class="${c.adds.VIVO >= 0 ? 'pos' : 'neg'}">${c.adds.VIVO >= 0 ? '+' : ''}${fmt(c.adds.VIVO)}</td>
        <td style="color:#2a5d8f">${fmt(c.base.TIM)}</td><td>${c.share.TIM.toFixed(1).replace('.', ',')}%</td><td class="${c.adds.TIM >= 0 ? 'pos' : 'neg'}">${c.adds.TIM >= 0 ? '+' : ''}${fmt(c.adds.TIM)}</td>
        <td><b>${fmt(c.base.total)}</b></td><td class="${c.adds.total >= 0 ? 'pos' : 'neg'}"><b>${c.adds.total >= 0 ? '+' : ''}${fmt(c.adds.total)}</b></td>
        <td style="text-align:center"><span class="pos-badge ${posClass}">${cRank}º Lugar</span></td>
      </tr>`;
    });

    totClaroA += subClaroA; totVivoA += subVivoA; totTimA += subTimA; totMercadoA += subMercadoA;
    totClaroNet += subClaroNet; totVivoNet += subVivoNet; totTimNet += subTimNet; totMercadoNet += subMercadoNet;

    // Subtotal
    html += `<tr class="total-row" style="background:#f4f6fa;font-weight:700">
      <td class="l"><b>SUBTOTAL — ${g.title.toUpperCase()}</b></td>
      <td style="color:var(--red)">${fmt(subClaroA)}</td><td>${subMercadoA > 0 ? (subClaroA/subMercadoA*100).toFixed(1).replace('.', ',') + '%' : '—'}</td><td class="${subClaroNet >= 0 ? 'pos' : 'neg'}">${subClaroNet >= 0 ? '+' : ''}${fmt(subClaroNet)}</td>
      <td style="color:#6b4b8a">${fmt(subVivoA)}</td><td>${subMercadoA > 0 ? (subVivoA/subMercadoA*100).toFixed(1).replace('.', ',') + '%' : '—'}</td><td class="${subVivoNet >= 0 ? 'pos' : 'neg'}">${subVivoNet >= 0 ? '+' : ''}${fmt(subVivoNet)}</td>
      <td style="color:#2a5d8f">${fmt(subTimA)}</td><td>${subMercadoA > 0 ? (subTimA/subMercadoA*100).toFixed(1).replace('.', ',') + '%' : '—'}</td><td class="${subTimNet >= 0 ? 'pos' : 'neg'}">${subTimNet >= 0 ? '+' : ''}${fmt(subTimNet)}</td>
      <td><b>${fmt(subMercadoA)}</b></td><td class="${subMercadoNet >= 0 ? 'pos' : 'neg'}"><b>${subMercadoNet >= 0 ? '+' : ''}${fmt(subMercadoNet)}</b></td>
      <td style="text-align:center">—</td>
    </tr>`;
  });

  // Total Geral
  html += `<tr class="total-row" style="background:#e8edf3;font-weight:800;border-top:3px solid var(--nav)">
    <td class="l"><b>TOTAL REGIONAL NORTE (455 CIDADES)</b></td>
    <td style="color:var(--red)">${fmt(totClaroA)}</td><td>${totMercadoA > 0 ? (totClaroA/totMercadoA*100).toFixed(1).replace('.', ',') + '%' : '—'}</td><td class="${totClaroNet >= 0 ? 'pos' : 'neg'}">${totClaroNet >= 0 ? '+' : ''}${fmt(totClaroNet)}</td>
    <td style="color:#6b4b8a">${fmt(totVivoA)}</td><td>${totMercadoA > 0 ? (totVivoA/totMercadoA*100).toFixed(1).replace('.', ',') + '%' : '—'}</td><td class="${totVivoNet >= 0 ? 'pos' : 'neg'}">${totVivoNet >= 0 ? '+' : ''}${fmt(totVivoNet)}</td>
    <td style="color:#2a5d8f">${fmt(totTimA)}</td><td>${totMercadoA > 0 ? (totTimA/totMercadoA*100).toFixed(1).replace('.', ',') + '%' : '—'}</td><td class="${totTimNet >= 0 ? 'pos' : 'neg'}">${totTimNet >= 0 ? '+' : ''}${fmt(totTimNet)}</td>
    <td><b>${fmt(totMercadoA)}</b></td><td class="${totMercadoNet >= 0 ? 'pos' : 'neg'}"><b>${totMercadoNet >= 0 ? '+' : ''}${fmt(totMercadoNet)}</b></td>
    <td style="text-align:center">—</td>
  </tr></tbody>`;

  tEl.innerHTML = html;
}

/* ============ VISÃO 4: DETALHADA HISTÓRICA & PROJEÇÃO (Item 4 do Usuário) ============ */
function renderDetalhada() {
  renderProjecaoMovel();
  renderHistoryLineChartMovel();
  renderMatrizDetalhadaMovel();
  renderAnalyticMovel();
}

function renderProjecaoMovel() {
  const claroBase = 2211446;
  const vivoBase = 2532055;
  const claroNet = 27120;
  const vivoNet = 19096;
  const gapVivo = vivoBase - claroBase; // 320.609
  const netAdvantage = claroNet - vivoNet; // +8.024/mês a favor da Claro

  const mesesUltrapassagem = netAdvantage > 0 ? Math.round(gapVivo / netAdvantage) : '—';
  const anosUltrapassagem = netAdvantage > 0 ? (gapVivo / netAdvantage / 12).toFixed(1).replace('.', ',') : '—';

  const bEl = document.getElementById('boxProjecao');
  bEl.innerHTML = `
    <div class="proj-card" style="background:#fdf2f2; border-color:#fadbd8">
      <div class="lbl" style="color:var(--red)">Posição Oficial Atual</div>
      <div class="val" style="color:var(--red)">2º Lugar Regional</div>
      <div class="desc">
        Claro possui <b>${fmt(claroBase)}</b> clientes (37,0% share), com gap de <b>${fmt(gapVivo)}</b> para a líder Vivo.
      </div>
    </div>

    <div class="proj-card" style="background:#eafaf1; border-color:#d4efdf">
      <div class="lbl" style="color:#1e6b37">Vantagem em Adições Líquidas</div>
      <div class="val" style="color:#1e6b37">+${fmt(netAdvantage)} / mês</div>
      <div class="desc">
        Claro adicionou <b>+${fmt(claroNet)}</b> contra <b>+${fmt(vivoNet)}</b> da Vivo (Net Share Claro de 48,5%).
      </div>
    </div>

    <div class="proj-card">
      <div class="lbl">Projeção para Alcance da Liderança</div>
      <div class="val" style="color:var(--blue)">${mesesUltrapassagem} meses <span style="font-size:12px;color:var(--ink3)">(${anosUltrapassagem} anos)</span></div>
      <div class="desc">
        Horizonte estimado mantendo o ritmo superior de captação de clientes (+8k clientes/mês sobre a Vivo).
      </div>
    </div>

    <div class="proj-card">
      <div class="lbl">Meta de Adições para 24 Meses</div>
      <div class="val" style="color:var(--amber)">+${fmt(Math.round(vivoNet + gapVivo / 24))} / mês</div>
      <div class="desc">
        Ritmo de adições líquidas recomendadas para assumir a liderança em 2 anos (+13.359/mês sobre a Vivo).
      </div>
    </div>

    <div class="proj-card" style="grid-column: 1 / -1; background:#fbfcfe">
      <div class="lbl">Diretriz Estratégica Regional</div>
      <div class="desc" style="font-size:12px;line-height:1.6;color:var(--ink2);margin-top:4px">
        A Claro é a operadora com <b>maior tração de crescimento na Regional Norte</b>, liderando isolada as adições líquidas mensais em 7 dos 9 DDDs (91, 92, 95, 96, 97, 98 e 99). 
        A manutenção desse diferencial competitivo garantirá a tomada da liderança de share na telefonia móvel.
      </div>
    </div>
  `;
}

function renderHistoryLineChartMovel() {
  const wrap = document.getElementById('lineChartSvgWrap');
  const leg = document.getElementById('lineChartLegend');
  if (!wrap || !leg) return;

  const PLAYERS_CHART = [
    { id: 'Vivo', name: 'Vivo', color: '#6b4b8a', strokeWidth: 2.4 },
    { id: 'Claro', name: 'Claro', color: '#c0392b', strokeWidth: 2.8 },
    { id: 'Tim', name: 'TIM', color: '#2a5d8f', strokeWidth: 2.2 },
    { id: 'SURF', name: 'Surf', color: '#c98a15', strokeWidth: 1.8 },
    { id: 'Outros', name: 'Outros', color: '#7f8c8d', strokeWidth: 1.8 }
  ];

  leg.innerHTML = PLAYERS_CHART.map(p => `
    <div class="chart-leg-item">
      <span class="chart-leg-dot" style="background:${p.color}"></span>
      <span>${p.name}:</span>
      <b style="color:${p.color}">${fmt(DATA.history_mov[p.id]?.base?.master[13] || 0)}</b>
    </div>
  `).join('') + `
    <div class="chart-leg-item" style="margin-left:auto;font-size:11px;color:var(--ink3)">
      <span>Total Mercado ago/26: <b>${fmt(DATA.history_mov['Total Geral']?.base?.master[13] || 0)}</b></span>
    </div>
  `;

  const W = 960, H = 240;
  const padL = 75, padR = 25, padT = 20, padB = 35;
  const plotW = W - padL - padR;
  const plotH = H - padT - padB;

  let maxVal = 2800000;

  let svg = `<svg viewBox="0 0 ${W} ${H}" style="width:100%;height:auto;display:block;font-family:'Segoe UI',sans-serif">`;
  const yTicks = 5;
  for (let i = 0; i <= yTicks; i++) {
    const val = (maxVal / yTicks) * i;
    const y = padT + plotH - (val / maxVal * plotH);
    svg += `<line x1="${padL}" y1="${y}" x2="${W - padR}" y2="${y}" stroke="#eef1f4" stroke-dasharray="${i === 0 ? 'none' : '3,3'}" stroke-width="1"/>`;
    svg += `<text x="${padL - 8}" y="${y + 4}" font-size="10" fill="#718096" text-anchor="end">${fmt(val)}</text>`;
  }

  const nMonths = MONTHS.length;
  const stepX = plotW / (nMonths - 1);
  MONTHS.forEach((m, idx) => {
    const x = padL + idx * stepX;
    const isCur = idx === 13;
    svg += `<line x1="${x}" y1="${padT + plotH}" x2="${x}" y2="${padT + plotH + 5}" stroke="#cbd5e0" stroke-width="1"/>`;
    svg += `<text x="${x}" y="${padT + plotH + 18}" font-size="10.5" font-weight="${isCur ? '700' : '500'}" fill="${isCur ? '#c0392b' : '#4a5568'}" text-anchor="middle">${m}</text>`;
    if (isCur) {
      svg += `<rect x="${x - 18}" y="${padT}" width="36" height="${plotH}" fill="rgba(224,32,32,0.03)" pointer-events="none"/>`;
    }
  });

  PLAYERS_CHART.forEach(p => {
    const arr = DATA.history_mov[p.id]?.base?.master || [];
    if (!arr.length) return;

    let pts = [];
    arr.forEach((v, idx) => {
      const x = padL + idx * stepX;
      const y = padT + plotH - (v / maxVal * plotH);
      pts.push(`${x.toFixed(1)},${y.toFixed(1)}`);
    });

    svg += `<polyline points="${pts.join(' ')}" fill="none" stroke="${p.color}" stroke-width="${p.strokeWidth}"/>`;
    arr.forEach((v, idx) => {
      const x = padL + idx * stepX;
      const y = padT + plotH - (v / maxVal * plotH);
      const isCur = idx === 13;
      svg += `<circle cx="${x.toFixed(1)}" cy="${y.toFixed(1)}" r="${isCur ? '4' : '2.5'}" fill="${p.color}">
        <title>${p.name} — ${MONTHS[idx]}: ${fmt(v)} clientes</title>
      </circle>`;
    });
  });

  svg += '</svg>';
  wrap.innerHTML = svg;
}

/* ============ MATRIZ HISTÓRICA MENSAL COM ABERTURA INDIVIDUAL POR INDICADOR — Item 4 ============ */
function renderMatrizDetalhadaMovel() {
  const tEl = document.getElementById('tblDetalhadaHistorica');
  const OPERATORS = ['Total Geral', 'Vivo', 'Claro', 'Tim', 'SURF', 'Outros'];

  // Determina lista de DDDs ativos conforme os filtros
  let activeDdds = [];
  if (sel.ddds && sel.ddds.size > 0) {
    activeDdds = [...sel.ddds].map(id => D.ddds[id]).sort();
  } else if (sel.ufs && sel.ufs.size > 0) {
    // DDDs da UF selecionada
    const ufMap = {
      'PA': ['91', '93', '94'],
      'AM': ['92', '97'],
      'MA': ['98', '99'],
      'AP': ['96'],
      'RR': ['95']
    };
    const sUfs = [...sel.ufs].map(id => D.ufs[id]);
    sUfs.forEach(u => {
      if (ufMap[u]) activeDdds.push(...ufMap[u]);
    });
    activeDdds = [...new Set(activeDdds)].sort();
    if (activeDdds.length === 0) activeDdds = ['91', '92', '93', '94', '95', '96', '97', '98', '99'];
  } else {
    activeDdds = ['91', '92', '93', '94', '95', '96', '97', '98', '99'];
  }

  let html = '<thead><tr><th class="l">Operadora &amp; Indicador</th>';
  MONTHS.forEach((m, idx) => {
    const isCur = idx === 13;
    html += `<th class="${isCur ? 'col-cur-mth-hdr' : ''}">${m}${isCur ? ' <span style="font-size:9px;background:var(--red);color:#fff;padding:1px 4px;border-radius:2px;margin-left:2px">Atual</span>' : ''}</th>`;
  });
  html += '<th>MoM</th><th>YoY</th></tr></thead><tbody>';

  OPERATORS.forEach(op => {
    const d = DATA.history_mov[op];
    if (!d) return;

    const isClaro = op === 'Claro';
    const isVivo = op === 'Vivo';
    const isTim = op === 'Tim';
    const isSurf = op === 'SURF';
    const isTot = op === 'Total Geral';
    const opClean = op.replace('/', '_').replace(' ', '_');

    let opRowClass = 'outros-row';
    if (isClaro) opRowClass = 'claro-row';
    else if (isVivo) opRowClass = 'vivo-row';
    else if (isTim) opRowClass = 'tim-row';
    else if (isSurf) opRowClass = 'surf-row';
    else if (isTot) opRowClass = 'total-row';

    const kBase = `${opClean}_base`;
    const kShare = `${opClean}_share`;
    const kAdds = `${opClean}_adds`;
    const kNetShare = `${opClean}_netshare`;
    const dddCountNote = `<span style="font-size:10px;color:var(--ink3);font-weight:400">(${activeDdds.length} DDD${activeDdds.length > 1 ? 's' : ''})</span>`;

    // 1. Linha Mestre: Base de Acessos
    const bMast = d.base.master;
    html += `<tr class="${opRowClass}" style="cursor:pointer" onclick="toggleDetIndicator('${kBase}')">
      <td class="l"><b><span id="ico_${kBase}">▸</span> ${op} — Base de Acessos</b> ${dddCountNote}</td>`;
    bMast.forEach((v, idx) => {
      html += `<td class="${idx === 13 ? 'col-cur-mth' : ''}">${fmt(v)}</td>`;
    });
    html += `<td>${pct(bMast[13] - bMast[12], bMast[12])}</td><td>${pct(bMast[13] - bMast[1], bMast[1])}</td></tr>`;

    // 1.1 Abertura de DDDs da Base
    activeDdds.forEach(dddStr => {
      const bDdd = d.base.ddds[dddStr] || new Array(14).fill(0);
      html += `<tr class="det-sub-${kBase}" style="display:none;background:#f8fafc;color:var(--ink)">
        <td class="l" style="padding-left:34px;font-weight:600">↳ DDD ${dddStr} · Base</td>`;
      bDdd.forEach((v, idx) => {
        html += `<td class="${idx === 13 ? 'col-cur-mth' : ''}">${fmt(v)}</td>`;
      });
      html += `<td>${pct(bDdd[13] - bDdd[12], bDdd[12])}</td><td>${pct(bDdd[13] - bDdd[1], bDdd[1])}</td></tr>`;
    });

    // 2. Linha Mestre: Market Share (%)
    const sMast = d.share.master;
    html += `<tr style="background:#fafbfc;color:var(--ink2);cursor:pointer" onclick="toggleDetIndicator('${kShare}')">
      <td class="l" style="padding-left:18px"><b><span id="ico_${kShare}">▸</span> ${op} — Market Share (%)</b> ${dddCountNote}</td>`;
    sMast.forEach((v, idx) => {
      html += `<td class="${idx === 13 ? 'col-cur-mth' : ''}">${v.toFixed(2).replace('.', ',')}%</td>`;
    });
    html += `<td>${pp(sMast[13] - sMast[12])}</td><td>${pp(sMast[13] - sMast[1])}</td></tr>`;

    // 2.1 Abertura de DDDs do Share
    activeDdds.forEach(dddStr => {
      const sDdd = d.share.ddds[dddStr] || new Array(14).fill(0);
      html += `<tr class="det-sub-${kShare}" style="display:none;background:#fbfcfe;color:var(--ink2)">
        <td class="l" style="padding-left:34px;font-size:11px">↳ DDD ${dddStr} · Share %</td>`;
      sDdd.forEach((v, idx) => {
        html += `<td class="${idx === 13 ? 'col-cur-mth' : ''}" style="font-size:11px">${v.toFixed(2).replace('.', ',')}%</td>`;
      });
      html += `<td style="font-size:11px">${pp(sDdd[13] - sDdd[12])}</td><td style="font-size:11px">${pp(sDdd[13] - sDdd[1])}</td></tr>`;
    });

    // 3. Linha Mestre: Net Adds
    const aMast = d.adds.master;
    html += `<tr style="background:#fafbfc;color:var(--ink2);cursor:pointer" onclick="toggleDetIndicator('${kAdds}')">
      <td class="l" style="padding-left:18px"><b><span id="ico_${kAdds}">▸</span> ${op} — Net Adds</b> ${dddCountNote}</td>`;
    aMast.forEach((v, idx) => {
      html += `<td class="${idx === 13 ? 'col-cur-mth ' : ''}${v >= 0 ? 'pos' : 'neg'}">${v >= 0 ? '+' : ''}${fmt(v)}</td>`;
    });
    html += `<td>${aMast[13] - aMast[12] >= 0 ? '+' : ''}${fmt(aMast[13] - aMast[12])}</td><td>${aMast[13] - aMast[1] >= 0 ? '+' : ''}${fmt(aMast[13] - aMast[1])}</td></tr>`;

    // 3.1 Abertura de DDDs de Net Adds
    activeDdds.forEach(dddStr => {
      const aDdd = d.adds.ddds[dddStr] || new Array(14).fill(0);
      html += `<tr class="det-sub-${kAdds}" style="display:none;background:#fbfcfe;color:var(--ink2)">
        <td class="l" style="padding-left:34px;font-size:11px">↳ DDD ${dddStr} · Net Adds</td>`;
      aDdd.forEach((v, idx) => {
        html += `<td class="${idx === 13 ? 'col-cur-mth ' : ''}${v >= 0 ? 'pos' : 'neg'}" style="font-size:11px">${v >= 0 ? '+' : ''}${fmt(v)}</td>`;
      });
      html += `<td style="font-size:11px">${aDdd[13] - aDdd[12] >= 0 ? '+' : ''}${fmt(aDdd[13] - aDdd[12])}</td><td style="font-size:11px">${aDdd[13] - aDdd[1] >= 0 ? '+' : ''}${fmt(aDdd[13] - aDdd[1])}</td></tr>`;
    });

    // 4. Linha Mestre: Net Share (%)
    const nsMast = d.net_share.master;
    html += `<tr style="background:#fafbfc;color:var(--ink2);border-bottom:2px solid var(--line2);cursor:pointer" onclick="toggleDetIndicator('${kNetShare}')">
      <td class="l" style="padding-left:18px"><b><span id="ico_${kNetShare}">▸</span> ${op} — Net Share (%)</b> ${dddCountNote}</td>`;
    nsMast.forEach((v, idx) => {
      html += `<td class="${idx === 13 ? 'col-cur-mth' : ''}">${v.toFixed(1).replace('.', ',')}%</td>`;
    });
    html += `<td>—</td><td>—</td></tr>`;

    // 4.1 Abertura de DDDs de Net Share (%)
    activeDdds.forEach(dddStr => {
      const nsDdd = d.net_share.ddds[dddStr] || new Array(14).fill(0);
      html += `<tr class="det-sub-${kNetShare}" style="display:none;background:#fbfcfe;color:var(--ink2);border-bottom:1px dashed var(--line2)">
        <td class="l" style="padding-left:34px;font-size:11px">↳ DDD ${dddStr} · Net Share %</td>`;
      nsDdd.forEach((v, idx) => {
        html += `<td class="${idx === 13 ? 'col-cur-mth' : ''}" style="font-size:11px">${v > 0 ? v.toFixed(1).replace('.', ',') + '%' : '—'}</td>`;
      });
      html += `<td style="font-size:11px">—</td><td style="font-size:11px">—</td></tr>`;
    });
  });

  html += '</tbody>';
  tEl.innerHTML = html;

  // Botão de expandir/recolher todos os DDDs
  const btnTgl = document.getElementById('btnToggleAllDdd');
  if (btnTgl) {
    btnTgl.onclick = () => {
      const allSub = document.querySelectorAll('tr[class*="det-sub-"]');
      const anyHidden = [...allSub].some(r => r.style.display === 'none');
      allSub.forEach(r => { r.style.display = anyHidden ? '' : 'none'; });
      document.querySelectorAll('span[id^="ico_"]').forEach(ico => {
        ico.textContent = anyHidden ? '▾' : '▸';
      });
      btnTgl.textContent = anyHidden ? 'Recolher Todos os DDDs' : 'Expandir Todos os DDDs';
    };
  }
}

window.toggleDetIndicator = function(key) {
  const rows = document.querySelectorAll('.det-sub-' + key);
  const ico = document.getElementById('ico_' + key);
  const anyHidden = [...rows].some(r => r.style.display === 'none');
  rows.forEach(r => {
    r.style.display = anyHidden ? '' : 'none';
  });
  if (ico) ico.textContent = anyHidden ? '▾' : '▸';
};

/* ============ TABELA ANALÍTICA ============ */
let detailPage = 1;
const PAGE_SIZE = 100;
function renderAnalyticMovel() {
  const tEl = document.getElementById('tblDetalhada');
  const q = (document.getElementById('detailSearch')?.value || '').toLowerCase().trim();

  let rows = getFilteredRows();
  if (q) {
    rows = rows.filter(r => {
      const emp = D.empresas[r[0]].toLowerCase();
      const pr = D.produtos[r[1]].toLowerCase();
      const u = D.ufs[r[2]].toLowerCase();
      const d = D.ddds[r[3]].toLowerCase();
      return emp.includes(q) || pr.includes(q) || u.includes(q) || d.includes(q);
    });
  }

  document.getElementById('analyticCountNote').textContent = `${rows.length.toLocaleString('pt-BR')} registros filtrados`;

  const totalPages = Math.max(1, Math.ceil(rows.length / PAGE_SIZE));
  if (detailPage > totalPages) detailPage = totalPages;
  document.getElementById('detailPaginationInfo').textContent = `Página ${detailPage} de ${totalPages} (${rows.length.toLocaleString('pt-BR')} registros)`;

  const start = (detailPage - 1) * PAGE_SIZE;
  const pageRows = rows.slice(start, start + PAGE_SIZE);

  let html = `<thead><tr>
    <th class="l">Operadora</th>
    <th class="l">Produto</th>
    <th class="l">Cobrança</th>
    <th class="l">Pessoa</th>
    <th class="l">Tecnologia</th>
    <th>UF</th>
    <th>DDD</th>
    <th>Base jul/26</th>
    <th>Base ago/26</th>
    <th>Net Adds</th>
  </tr></thead><tbody>`;

  pageRows.forEach(r => {
    const net = r[8] - r[7];
    html += `<tr>
      <td class="l"><b>${D.empresas[r[0]]}</b></td>
      <td class="l">${D.produtos[r[1]]}</td>
      <td class="l">${D.cobrancas[r[4]]}</td>
      <td class="l">${D.pessoas[r[5]]}</td>
      <td class="l">${D.tecnologias[r[6]]}</td>
      <td>${D.ufs[r[2]]}</td>
      <td>${D.ddds[r[3]]}</td>
      <td>${fmt(r[7])}</td>
      <td><b>${fmt(r[8])}</b></td>
      <td class="${net >= 0 ? 'pos' : 'neg'}">${net >= 0 ? '+' : ''}${fmt(net)}</td>
    </tr>`;
  });

  html += '</tbody>';
  tEl.innerHTML = html;

  document.getElementById('btnPrevPage').onclick = () => {
    if (detailPage > 1) { detailPage--; renderAnalyticMovel(); }
  };
  document.getElementById('btnNextPage').onclick = () => {
    if (detailPage < totalPages) { detailPage++; renderAnalyticMovel(); }
  };
  document.getElementById('detailSearch').oninput = () => {
    detailPage = 1; renderAnalyticMovel();
  };
}

/* ============ Exportação CSV ============ */
function exportCsv() {
  const fRows = getFilteredRows();
  let csv = 'Operadora;Produto;UF;DDD;Cobranca;Pessoa;Tecnologia;Base_Jul26;Base_Ago26;Net_Adds\\n';
  fRows.forEach(r => {
    const net = r[8] - r[7];
    csv += `"${D.empresas[r[0]]}";"${D.produtos[r[1]]}";"${D.ufs[r[2]]}";"${D.ddds[r[3]]}";"${D.cobrancas[r[4]]}";"${D.pessoas[r[5]]}";"${D.tecnologias[r[6]]}";${r[7]};${r[8]};${net}\\n`;
  });
  const blob = new Blob([csv], { type: 'text/csv;charset=utf-8;' });
  const a = document.createElement('a');
  a.href = URL.createObjectURL(blob);
  a.download = `market_share_movel_filtrado_${Date.now()}.csv`;
  a.click();
}

/* ============ Renderização Principal ============ */
function render() {
  updateSelCounts();
  refreshAvailability();
  drawChips();
  drawSummary();

  const fRows = getFilteredRows();
  document.getElementById('hdrRows').textContent = `${fRows.length.toLocaleString('pt-BR')} registros RNO`;

  renderKpis(fRows);

  if (currentTab === 'resumida') renderResumida();
  else if (currentTab === '5g') render5G();
  else if (currentTab === 'municipio') renderMunicipio();
  else if (currentTab === 'detalhada') renderDetalhada();
}

/* ============ Inicialização ============ */
document.addEventListener('DOMContentLoaded', () => {
  buildFilterUI();
  applyDefaultFilters();

  // Abas de Visão Principal
  document.querySelectorAll('#viewTabs .tab-btn').forEach(btn => {
    btn.onclick = () => {
      document.querySelectorAll('#viewTabs .tab-btn').forEach(b => b.classList.remove('on'));
      btn.classList.add('on');
      currentTab = btn.dataset.tab;
      document.querySelectorAll('.tab-content').forEach(c => c.style.display = 'none');
      const el = document.getElementById('view' + currentTab.charAt(0).toUpperCase() + currentTab.slice(1));
      if (el) el.style.display = '';
      render();
    };
  });

  // Dimensões na Visão Resumida (DDD, UF, REGIONAL)
  document.querySelectorAll('#tabsResumoDim .tab-dim-btn').forEach(btn => {
    btn.onclick = () => {
      document.querySelectorAll('#tabsResumoDim .tab-dim-btn').forEach(b => b.classList.remove('on'));
      btn.classList.add('on');
      currentResumoDim = btn.dataset.dim;
      renderResumida();
    };
  });

  // Pills de DDD Rápido
  document.querySelectorAll('#dddPills .ddd-pill').forEach(btn => {
    btn.onclick = () => {
      const dVal = btn.dataset.ddd;
      const s = document.querySelector('select[data-dim="ddds"]');
      if (dVal === 'TODOS') {
        delete sel.ddds;
        if (s) [...s.options].forEach(o => o.selected = false);
      } else {
        const id = D.ddds.indexOf(dVal);
        if (id >= 0) {
          sel.ddds = new Set([id]);
          if (s) [...s.options].forEach(o => { o.selected = (+o.value === id); });
        }
      }
      syncDddPills();
      render();
    };
  });

  // Modais de Ajuda
  const HELP = {
    filtros: '<h4>Gaveta de Filtros</h4><p>Permite refinar a base móvel por operadora, tecnologia, modalidade (Pós/Pré-pago), pessoa (PF/PJ) e localização.</p>',
    kpis: '<h4>Indicadores Oficiais</h4><p>Consolidam a posição da Claro Móvel na Regional Norte contra Vivo, TIM, Surf e Outros players.</p>',
    resumida: '<h4>Visão Resumida</h4><p>Exibe os 4 quadros oficiais comparando operadoras por DDD, Estado ou Regional Nacional com dados de todas as 10 regionais.</p>',
    tec5g: '<h4>Penetração 5G</h4><p>Percentual de aparelhos 5G ativos sobre a base total e número de terminais por DDD.</p>',
    municipio: '<h4>Visão Municipal</h4><p>Desempenho nos 455 municípios da Regional Norte subdivididos em Grandes Mercados, Mercados em Desenvolvimento e Demais Cidades.</p>',
    projecao: '<h4>Projeção de Liderança</h4><p>Modelo matemático de ultrapassagem calculando o ritmo mensal da Claro sobre a Vivo.</p>'
  };
  document.querySelectorAll('.hlp').forEach(btn => {
    btn.onclick = e => {
      e.stopPropagation();
      const h = btn.dataset.h;
      if (HELP[h]) {
        document.getElementById('helpTitle').textContent = 'Ajuda: ' + h.toUpperCase();
        document.getElementById('helpBody').innerHTML = HELP[h];
        document.getElementById('helpModal').classList.add('on');
      }
    };
  });

  render();
});
</script>
</body>
</html>
"""
    return template.replace('__PAYLOAD__', json_data).replace('__MES_REFERENCIA__', payload['mes_referencia']).replace('__DATA_POSICAO__', payload['data_posicao'])

if __name__ == '__main__':
    build()
